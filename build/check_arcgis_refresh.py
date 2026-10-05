"""Golden check for the in-app ArcGIS layer refresh (v0.46.0): the worker that
runs in ArcGIS Pro's Python (scripts/arcgis_worker/export_layers.py), the
session runner (arcgis_pro), the refresh/probe orchestration (arcgis_refresh),
the library manifest + drop identity (arcgis_layers), and the GUI wiring.

ArcPy exists only inside ArcGIS Pro, so every run here drives the SHIPPED worker
script against the stand-in ArcPy in arcgis_selftest — in-process for the
orchestration cases, and once as a REAL subprocess (this venv's python.exe with
a stand-in `arcpy.py` beside a copy of the worker) through the real launcher,
environment and event-file protocol. Offline; openpyxl only.

    build\\.venv\\Scripts\\python.exe build\\check_arcgis_refresh.py
"""
import ast
import contextlib
import functools
import io
import json
import logging
import os
import queue
import shutil
import sys
import tempfile
import threading
import urllib.request
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(ROOT / "scripts"), str(ROOT)]

import arcgis_layers  # noqa: E402
import arcgis_pro  # noqa: E402
import arcgis_refresh as ar  # noqa: E402
import arcgis_selftest as ast_  # noqa: E402
import clean_road_layers as crl  # noqa: E402
import paths  # noqa: E402
from events import Events  # noqa: E402

_fail = []


def check(name, cond, detail=""):
    print(f"  [{'OK ' if cond else 'FAIL'}] {name}" + (f" — {detail}" if detail and not cond else ""))
    if not cond:
        _fail.append(name)


@contextlib.contextmanager
def _patch(obj, name, value):
    old = getattr(obj, name)
    setattr(obj, name, value)
    try:
        yield
    finally:
        setattr(obj, name, old)


@contextlib.contextmanager
def _library():
    """A temp library seeded with a manual drop; ARCGIS_LAYERS_ROOT points at it."""
    tmp = Path(tempfile.mkdtemp(prefix="agrefresh_"))
    lib = tmp / "arcgis_layers"
    ast_.seed_manual_library(lib)
    with _patch(paths, "ARCGIS_LAYERS_ROOT", lib):
        try:
            yield lib
        finally:
            shutil.rmtree(tmp, ignore_errors=True)


def _refresh(spec, layers, **kw):
    fake = ast_.make_fake_arcpy(spec)
    logs = []
    res = ar.refresh(layers, events=Events(on_log=logs.append), python="fake-python.exe",
                     launcher=ast_.inprocess_launcher(fake), **kw)
    return res, {r["name"]: r for r in res.layers}, logs


# --------------------------------------------------------------------------- #
def test_contract():
    print("the library contract — 40 layers, the worker script, its Python floor:")
    names = [n for n, _id in ar.LIBRARY_LAYERS]
    check("LIBRARY_LAYERS is exactly clean_road_layers.EXPECTED_LAYERS",
          sorted(names) == sorted(crl.EXPECTED_LAYERS) and len(names) == 40)
    ids = [i for _n, i in ar.LIBRARY_LAYERS]
    check("every fallback layer id is unique", len(set(ids)) == len(ids))
    check("the probe layer is a library layer", ar.PROBE_LAYER in ar.LAYER_IDS)
    script = arcgis_pro.worker_script()
    check("the worker script ships beside the modules", script.is_file(), str(script))
    src = script.read_text(encoding="utf-8")
    try:
        tree = ast.parse(src, feature_version=(3, 9))
        parsed = True
    except SyntaxError as e:
        tree, parsed = None, False
        print(f"     {e}")
    check("the worker parses as Python 3.9 (ArcGIS Pro 3.0's floor)", parsed)
    allowed = {"argparse", "json", "os", "re", "sys", "threading", "time", "traceback",
               "urllib", "urllib.parse", "urllib.request", "arcpy"}
    used = set()
    for node in ast.walk(tree) if tree else ():
        if isinstance(node, ast.Import):
            used |= {a.name for a in node.names}
        elif isinstance(node, ast.ImportFrom) and node.module:
            used.add(node.module)
    check("the worker imports only the standard library + ArcPy",
          used <= allowed, f"unexpected: {sorted(used - allowed)}")
    check("the worker never writes the sign-in token into an event",
          "token=token" not in src and '"token":' not in src.replace('"token": token["token"]', ""))


def test_scenario():
    print("the shared scenario (also the frozen exe's self-test):")
    tmp = Path(tempfile.mkdtemp(prefix="agscenario_"))
    try:
        ast_.run(tmp, emit=lambda line: print(f"     {line}"))
        check("arcgis_selftest.run passes", True)
    except AssertionError as e:
        check("arcgis_selftest.run passes", False, str(e))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_truncated_export_is_refused():
    print("a truncated export never reaches the library:")
    with _library() as lib:
        before = (lib / "11_SHS Tolls.xlsx").read_bytes()
        spec = ast_.sample_spec()
        spec["layers"]["138"]["truncate"] = 2
        res, by, _logs = _refresh(spec, ["SHS Tolls"])
        check("the layer fails as truncated",
              by["SHS Tolls"]["status"] == "failed"
              and "truncated" in (by["SHS Tolls"].get("message") or ""),
              str(by["SHS Tolls"]))
        check("the run reads failed", res.status == "failed", res.status)
        check("the library's file is untouched",
              (lib / "11_SHS Tolls.xlsx").read_bytes() == before)
        check("the manifest still says 3 rows", crl.read_index(lib)["SHS Tolls"]["rows"] == 3)
        check("nothing was backed up", not (lib / "_previous").exists())


def test_cancel():
    print("Cancel stops the worker; finished layers stay, the rest are untouched:")
    with _library() as lib:
        cancel = threading.Event()
        spec = ast_.sample_spec()
        spec["layers"]["74"]["delay"] = 30          # City stalls in Table To Excel
        fake = ast_.make_fake_arcpy(spec)

        def on_progress(p):
            if p.get("done") == 1:                  # SHS Tolls is in: cancel
                cancel.set()

        res = ar.refresh(["SHS Tolls", "City"], events=Events(), cancel_event=cancel,
                         on_progress=on_progress, python="fake-python.exe",
                         launcher=ast_.inprocess_launcher(fake))
        by = {r["name"]: r for r in res.layers}
        check("the run reads cancelled", res.status == "cancelled", res.status)
        check("the finished layer was kept", by["SHS Tolls"]["status"] == "exported")
        check("the stopped layer reads cancelled", by["City"]["status"] == "cancelled",
              str(by["City"]))
        check("the stopped layer never reached the library", not (lib / "33_City.xlsx").exists())
        check("staging is empty after a cancel", not any(ar.staging_dir().glob("*")))


def test_catalog_by_name():
    print("signed in, layers are matched by NAME from the service's own list:")

    class _Resp(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    service_json = {"layers": [{"id": 999, "name": "SHS Tolls"},
                               {"id": 74, "name": "City"}],
                    "tables": []}
    seen = {}

    def fake_urlopen(req, timeout=None):
        seen["url"] = req.full_url
        seen["body"] = req.data.decode()
        return _Resp(json.dumps(service_json).encode())

    with _library() as lib, _patch(urllib.request, "urlopen", fake_urlopen):
        spec = ast_.sample_spec(signed_in=True)
        spec["layers"]["999"] = spec["layers"].pop("138")     # the service moved it
        res, by, logs = _refresh(spec, ["SHS Tolls", "City", "Route Direction"])
        check("a moved layer is found by name", by["SHS Tolls"]["status"] == "exported",
              str(by["SHS Tolls"]))
        check("a layer the service doesn't list fails by name",
              by["Route Direction"]["status"] == "failed"
              and "no layer named" in (by["Route Direction"].get("message") or ""))
        check("the token goes in the POST body, never the URL",
              "token=" not in seen.get("url", "") and "token=" in seen.get("body", ""))
        check("no log line carries the token", not any("fake-token" in m for m in logs))
        check("the manifest records the URL the layer was read from",
              crl.read_index(lib)["SHS Tolls"]["source"].endswith("/FeatureServer/999"))


def test_index_write_failure_rolls_back():
    print("a manifest that can't be written undoes that layer's swap:")
    with _library() as lib:
        before = (lib / "11_SHS Tolls.xlsx").read_bytes()

        def locked(*_a, **_k):
            raise PermissionError("[WinError 32] 00_INDEX.xlsx is open in Excel")

        logging.disable(logging.CRITICAL)        # the expected failure logs a traceback
        try:
            with _patch(ar, "write_index", locked):
                res, by, _logs = _refresh(ast_.sample_spec(), ["SHS Tolls"])
        finally:
            logging.disable(logging.NOTSET)
        check("the layer fails with the Excel hint",
              by["SHS Tolls"]["status"] == "failed"
              and "close it in Excel" in (by["SHS Tolls"].get("message") or ""),
              str(by["SHS Tolls"]))
        check("the library file is the one from before",
              (lib / "11_SHS Tolls.xlsx").read_bytes() == before)
        check("the manifest is the one from before", crl.read_index(lib)["SHS Tolls"]["rows"] == 3)


def test_not_signed_in_hint():
    print("a refresh that reads nothing while signed out says so plainly:")
    with _library():
        spec = ast_.sample_spec()
        spec["layers"] = {}                                    # every read refused
        res, _by, _logs = _refresh(spec, ["SHS Tolls"])
        check("the run reads failed", res.status == "failed", res.status)
        check("the message says ArcGIS Pro is not signed in",
              "not signed in" in res.message and "Sign me in automatically" in res.message,
              res.message)
    hint = ar.signin_hint({"signed_in": True, "portal": "https://www.arcgis.com/"})
    check("a different active portal is named", "not the TSMIS portal" in hint, hint)
    check("a signed-in TSMIS portal needs no hint",
          ar.signin_hint({"signed_in": True,
                          "portal": "https://rhapps-prod.dot.ca.gov/portal/"}) == "")


def test_real_subprocess():
    print("the REAL launcher: a separate python.exe, the event file, exit + crash + stall:")
    tmp = Path(tempfile.mkdtemp(prefix="agproc_"))
    worker_dir = tmp / "worker"
    worker_dir.mkdir()
    shutil.copy2(arcgis_pro.worker_script(), worker_dir / arcgis_pro.WORKER_SCRIPT)
    (worker_dir / "arcpy.py").write_text(ast_.FAKE_ARCPY_SOURCE, encoding="utf-8")
    spec_path = tmp / "fake_arcpy.json"
    lib = tmp / "arcgis_layers"
    ast_.seed_manual_library(lib)
    script = worker_dir / arcgis_pro.WORKER_SCRIPT
    os.environ["TSMIS_FAKE_ARCPY"] = str(spec_path)     # passes the clean env through
    try:
        with _patch(paths, "ARCGIS_LAYERS_ROOT", lib), \
                _patch(arcgis_pro, "worker_script", lambda: script):
            spec_path.write_text(json.dumps(ast_.sample_spec()), encoding="utf-8")
            res = ar.refresh(["SHS Tolls", "City"], events=Events(), python=sys.executable)
            by = {r["name"]: r for r in res.layers}
            check("a real worker process exports both layers", res.status == "ok",
                  f"{res.status}: {res.message}")
            check("its rows reach the library",
                  by.get("City", {}).get("rows") == 40 and (lib / "33_City.xlsx").is_file())
            out = (ar.session_dir() / "worker-output.txt")
            check("the worker's console output is kept for diagnosis", out.is_file())

            spec_path.write_text(json.dumps(ast_.sample_spec(crash=True)), encoding="utf-8")
            res = ar.refresh(["SHS Tolls"], events=Events(), python=sys.executable)
            check("a worker that dies is a failed run, not a hang",
                  res.status == "failed", f"{res.status}: {res.message}")

            spec_path.write_text(json.dumps(ast_.sample_spec(import_delay=30)), encoding="utf-8")
            outcome = arcgis_pro.run_worker(
                ar.session_dir(), {"protocol": 1, "mode": "probe", "service": "x",
                                   "probe_layer": {"name": "SHS Tolls", "id": 138,
                                                   "out": str(tmp / "p.xlsx")}},
                lambda _ev: None, None, sys.executable, start_limit=2)
            check("a worker that never starts the run is stopped at the start limit",
                  "did not start" in outcome.timed_out, outcome.timed_out)
    finally:
        os.environ.pop("TSMIS_FAKE_ARCPY", None)
        shutil.rmtree(tmp, ignore_errors=True)


def test_drop_dates():
    print("the drop's dates — every layer's own export time, the oldest wins:")
    from openpyxl import Workbook
    tmp = Path(tempfile.mkdtemp(prefix="agdates_"))
    try:
        for name in ("01_SHS Tolls.xlsx", "02_City.xlsx"):
            wb = Workbook()
            wb.active.append(["OBJECTID"])
            wb.save(tmp / name)
        rows = {"SHS Tolls": {"file": "01_SHS Tolls.xlsx", "rows": 0, "fields": 1,
                              "exported_at": "2026-09-30T08:00:00"},
                "City": {"file": "02_City.xlsx", "rows": 0, "fields": 1,
                         "exported_at": "2026-10-02T14:05:00"}}
        arcgis_layers.write_index(tmp, rows)
        drop = arcgis_layers.drop_info(tmp)
        check("the drop dates to its OLDEST layer",
              drop["exported"] == "2026-09-30" and drop["exported_source"] == "index", str(drop))
        check("a drop exported over several days says so",
              drop["mixed"] is True and drop["newest_at"] == "2026-10-02T14:05:00")
        check("a build reading only City may go as of City's own date",
              arcgis_layers.consistent_asof(["City"], drop) == "2026-10-02")
        check("a build reading both is held to the oldest",
              arcgis_layers.consistent_asof(["City", "SHS Tolls"], drop) == "2026-09-30")
        check("a layer with no recorded time falls back to the drop's date",
              arcgis_layers.consistent_asof(["Route Direction"], drop) == "2026-09-30")
        full, created = arcgis_layers.read_index_full(tmp)
        check("the manifest's own stamp is the oldest layer's time",
              created is not None and created.isoformat() == "2026-09-30T08:00:00",
              str(created))
        check("the six classic columns still read with the builds' reader",
              crl.read_index(tmp)["City"]["file"] == "02_City.xlsx")
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_row_count():
    print("row counting reads the FIRST sheet's XML, exactly:")
    from openpyxl import Workbook
    tmp = Path(tempfile.mkdtemp(prefix="agrows_"))
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["a", "b"])
        for i in range(2500):
            ws.append([i, f"text {i}"])
        other = wb.create_sheet("Other")
        for i in range(7):
            other.append([i])
        wb.save(tmp / "w.xlsx")
        check("2,501 rows in the first sheet (header + 2,500)",
              arcgis_layers.count_sheet_rows(tmp / "w.xlsx") == 2501)
        info = arcgis_layers.inspect_export(tmp / "w.xlsx", server_rows=2500)
        check("inspect_export: 2,500 data rows, 2 columns",
              info["data_rows"] == 2500 and info["fields"] == 2)
        try:
            arcgis_layers.inspect_export(tmp / "w.xlsx", server_rows=2501)
            refused = False
        except ValueError:
            refused = True
        check("fewer rows than the server counted is refused", refused)
        (tmp / "bad.xlsx").write_bytes(b"not a workbook")
        try:
            arcgis_layers.inspect_export(tmp / "bad.xlsx")
            bad = False
        except ValueError:
            bad = True
        check("an unreadable file is refused as ValueError", bad)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def test_gui_wiring():
    print("the GUI wiring — endpoints, the queue job, the worker's messages:")
    import gui_api
    import gui_arcgis_layers_api as gal
    import gui_matrix
    import gui_worker_arcgis as gwa

    for name in ("arcgis_layers_info", "refresh_arcgis_layers", "check_arcgis_pro",
                 "choose_arcgis_python", "clear_arcgis_python"):
        check(f"GuiApi exposes {name}", callable(getattr(gui_api.GuiApi, name, None)))
    check("GuiApi carries the Layers mixin", gal.GuiArcgisLayersMixin in gui_api.GuiApi.__mro__)

    started = []

    class _StubWorker:
        def __init__(self, *args):
            started.append(args)

        def start(self):
            pass

    class _Host(gal.GuiArcgisLayersMixin):
        def __init__(self):
            self._lock = threading.RLock()
            self._matrix = None
            self.cancel_event = threading.Event()
            self.logs = []

        def _emit_log(self, m):
            self.logs.append(m)

        def _set_dot(self, *_a):
            pass

        def _emit(self, _e):
            pass

        def _gated_queue(self):
            return "Q"

    host = _Host()
    job = gui_matrix.GuiMatrixMixin._make_job.__get__(_JobFactory())(
        "arcgis_refresh", "cell", "Refresh 2 ArcGIS layers", which="arcgis",
        layers=["SHS Tolls", "City"])
    check("a refresh job carries its layers", job["layers"] == ["SHS Tolls", "City"])
    with _patch(gal, "ArcgisLayerRefreshWorker", _StubWorker):
        ok = host._dispatch_arcgis_refresh_job(job)
    check("dispatch starts the refresh worker with the layers",
          ok and started and started[-1][0] == ["SHS Tolls", "City"])
    check("dispatch sets the layers phase and total",
          host._matrix == {"phase": "layers", "row": None, "cell": None, "done": 0, "total": 2})

    # The real worker thread over the in-process stand-in: its message stream.
    with _library():
        fake = ast_.make_fake_arcpy(ast_.sample_spec())
        wrapped = functools.partial(ar.refresh, python="fake-python.exe",
                                    launcher=ast_.inprocess_launcher(fake))
        q = queue.Queue()
        with _patch(ar, "refresh", wrapped):
            w = gwa.ArcgisLayerRefreshWorker(["SHS Tolls", "Route Direction"], q,
                                             threading.Event())
            w.start()
            w.join(timeout=60)
        msgs = []
        while not q.empty():
            msgs.append(q.get())
        kinds = [k for k, _p in msgs]
        done = [p for k, p in msgs if k == "matrix_done"]
        check("exactly one matrix_done ends the job", kinds.count("matrix_done") == 1
              and kinds[-1] == "matrix_done")
        check("matrix_done names the job and counts the layers",
              done and done[0]["label"] == "Layer refresh" and done[0]["total"] == 2
              and done[0]["succeeded"] == 1 and done[0]["failed"] == 1, str(done))
        check("progress goes out as matrix_cell with the layer as the row",
              any(k == "matrix_cell" and p.get("row") == "SHS Tolls" for k, p in msgs))


class _JobFactory:
    """Just enough of GuiApi for _make_job (the shared id sequence)."""

    class _Coord:
        n = 0

        def next_seq(self):
            self.n += 1
            return self.n

    _coord = _Coord()


def main():
    print("=== ArcGIS layer refresh (in-app export) ===")
    test_contract()
    test_scenario()
    test_truncated_export_is_refused()
    test_cancel()
    test_catalog_by_name()
    test_index_write_failure_rolls_back()
    test_not_signed_in_hint()
    test_real_subprocess()
    test_drop_dates()
    test_row_count()
    test_gui_wiring()
    print()
    if _fail:
        print(f"FAILED: {len(_fail)} check(s): {_fail}")
        return 1
    print("ALL ARCGIS-REFRESH CHECKS PASSED")
    return 0


if __name__ == "__main__":
    sys.exit(main())
