"""Refresh the ArcGIS layer library from the app (v0.46.0).

Until v0.46.0 the owner exported the 40 TSMIS layers by hand in ArcGIS Pro and
copied the files into `arcgis_layers/`. This module does that refresh itself:
ArcGIS Pro's own Python runs `arcgis_worker/export_layers.py` (see
`arcgis_pro`), which writes one Table To Excel export per layer into a staging
folder, and every finished layer is checked and swapped into the library while
the next one exports:

  * CHECKED -- the file opens, its header is readable, and its data rows are
    counted from the sheet itself; fewer rows than ArcGIS counted on the server
    is a truncated export and the layer is not swapped in (a few MORE is the
    measured count/export race the library already tolerates);
  * SWAPPED -- the layer keeps its library file name (so a refreshed library
    reads exactly like a manual one), the file it replaces moves to
    `arcgis_layers/_previous/` (one generation, cleared by the next refresh
    that replaces something), and 00_INDEX.xlsx is rewritten with the layer's
    row/field counts, its FeatureServer URL and two new columns -- Exported At
    (when that layer's export started) and Exported By. The existing readers
    only read the first six columns, so manual and app drops stay
    interchangeable; if the manifest can't be written the swap is undone;
  * RECORDED -- `arcgis_layers/_refresh/last_run.json` keeps every layer's
    outcome for the Layers tab, rewritten after each layer.

The run happens on the matrix queue, so no build or comparison reads the
library while it changes. Console-free.
"""
from __future__ import annotations

import json
import logging
import os
import shutil
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import arcgis_pro
import clean_road_layers as crl
from arcgis_layers import inspect_export, read_header, read_index_full, write_index
from arcgis_pro import SessionOutcome
from events import Events

log = logging.getLogger("tsmis.arcgis_refresh")

# The TSMIS feature service the manual exports came from (prod SSOR, default
# version) -- the same `lrs_tsmis` service the site's SSOR prod reports query.
SERVICE_DEFAULT = ("https://rhapps-prod.dot.ca.gov/server/rest/services/"
                   "TSMIS/lrs_tsmis/FeatureServer")

# The 40-layer library in its file order (the NN_ prefix of a fresh library),
# with each layer's id on the default service as the 2026-08-19 manual export
# recorded them. The worker matches layers by NAME from the service's own layer
# list when it can; these ids are the fallback. Same names as
# clean_road_layers.EXPECTED_LAYERS (asserted by check_arcgis_refresh).
LIBRARY_LAYERS = (
    ("Traffic Volume Ramps", 157), ("SHS Route Break", 133),
    ("SHS Ramp Pt", 132), ("SHS Landmark", 123),
    ("IM Intersection Point", 0), ("Equation Points", 305),
    ("SHS Travel Way R", 140), ("SHS Travel Way L", 139),
    ("Traffic Volume Segments", 153), ("Terrain Type", 11),
    ("SHS Tolls", 138), ("SHS Surface Type R", 137),
    ("SHS Surface Type L", 136), ("SHS Special Feature R", 135),
    ("SHS Special Feature L", 134), ("SHS Ramp", 131),
    ("SHS Population", 130), ("SHS O Shld Width R", 129),
    ("SHS O Shld Width L", 128), ("SHS Non Add Mileage", 125),
    ("SHS Median", 124), ("SHS Inv Network Date", 122),
    ("SHS I Shld Width R", 121), ("SHS I Shld Width L", 120),
    ("SHS Highway Group", 116), ("SHS Forest HWY", 115),
    ("SHS District", 114), ("SHS Design Speed", 113),
    ("SHS Curb Landscape", 112), ("SHS Barrier", 110),
    ("SHS Access Control", 109), ("County Code", 85), ("City", 74),
    ("IM Complex Intersection Cross Reference", 212),
    ("IM Complex Intersection Influence Segments", 211),
    ("IM Intersection Approach Detail", 150),
    ("IM Intersection Approach Segments", 146),
    ("IM Intersection Detail", 151), ("IM Intersection Route Table", 147),
    ("Route Direction", 304),
)
LAYER_IDS = dict(LIBRARY_LAYERS)
LAYER_ORDER = {name: i for i, (name, _id) in enumerate(LIBRARY_LAYERS, start=1)}

# A 27-row layer: the ArcGIS Pro check reads and exports it in seconds.
PROBE_LAYER = "SHS Tolls"
PROBE_LAYER_LIMIT_S = 15 * 60

_LAST_RUN = "last_run.json"
_LAST_PROBE = "last_probe.json"


# --------------------------------------------------------------------------- #
# places
# --------------------------------------------------------------------------- #
def library_root():
    return Path(crl.root())


def work_dir():
    """App-owned working folder inside the library (sub-folders never count
    toward the library's content fingerprint)."""
    return library_root() / "_refresh"


def staging_dir():
    return work_dir() / "staging"


def session_dir():
    return work_dir() / "session"


def backup_dir():
    return library_root() / "_previous"


def _now_local():
    return datetime.now().replace(microsecond=0)


def _local_iso(epoch):
    return datetime.fromtimestamp(epoch).replace(microsecond=0).isoformat()


def _atomic_json(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"{path.name}.tmp-{os.getpid()}")
    tmp.write_text(json.dumps(data, indent=2, default=str), encoding="utf-8")
    os.replace(tmp, path)


def _read_json(path):
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else None
    except (OSError, ValueError):      # silent-ok: no record yet / a damaged one reads as none
        return None


def last_run():
    return _read_json(work_dir() / _LAST_RUN)


def last_probe():
    return _read_json(work_dir() / _LAST_PROBE)


def _clear_files(folder):
    """Remove the plain files directly inside an app-owned folder."""
    folder = Path(folder)
    if not folder.is_dir():
        return
    for p in folder.iterdir():
        try:
            if p.is_file() and not p.is_symlink():
                p.unlink()
        except OSError as e:
            log.info("arcgis_refresh: could not remove %s (%s)", p, e)


# --------------------------------------------------------------------------- #
# the refresh
# --------------------------------------------------------------------------- #
@dataclass
class RefreshResult:
    status: str                    # ok | partial | failed | cancelled | error
    message: str
    layers: list = field(default_factory=list)
    facts: dict = field(default_factory=dict)


def layer_users():
    """{layer: [report label, ...]}: which Reports-vs-layers builds read each
    layer (none yet = staged for a later build)."""
    import arcgis_reports

    users = {name: [] for name, _id in LIBRARY_LAYERS}
    for key in arcgis_reports.KEYS:
        if not arcgis_reports.can_build(key):
            continue
        s = arcgis_reports.spec(key)
        for layer in getattr(s.build, "REQUIRED_LAYERS", ()):
            if layer in users and s.label not in users[layer]:
                users[layer].append(s.label)
    return users


def _target_name(layer, inventory):
    present = inventory["present"].get(layer)
    if present is not None:
        return present.name
    return f"{LAYER_ORDER.get(layer, 99):02d}_{layer}.xlsx"


def _app_version():
    try:
        from version import __version__
        return __version__
    except ImportError:          # silent-ok: provenance text only
        return "?"


class _Refresher:
    """One refresh run's state, fed by the worker's events."""

    def __init__(self, layers, events, on_progress, service):
        self.lib = library_root()
        self.layers = layers
        self.events = events
        self.on_progress = on_progress or (lambda _p: None)
        self.service = service
        self.inventory = crl.inventory(self.lib)
        self.index, created = read_index_full(self.lib)
        legacy_at = created.isoformat() if created else None
        for row in self.index.values():          # a manual manifest: one stamp for all
            row["exported_at"] = row.get("exported_at") or legacy_at
        self.results = {name: {"name": name, "status": "queued"} for name in layers}
        self.facts = {}
        self.current = None
        self.started = time.time()
        self.layer_started = None
        self.backed_up = False
        self.done_count = 0

    # -- progress ------------------------------------------------------- #
    def progress(self, message=None):
        now = time.time()
        self.on_progress({
            "phase": "layers", "layer": self.current,
            "index": self.done_count + (1 if self.current else 0),
            "done": self.done_count, "total": len(self.layers),
            "elapsed_s": round(now - self.started, 1),
            "layer_elapsed_s": (round(now - self.layer_started, 1)
                                if self.layer_started else None),
            "message": message})

    def record(self):
        _atomic_json(work_dir() / _LAST_RUN, self.summary("running"))

    def summary(self, status, message=""):
        return {"status": status, "message": message, "service": self.service,
                "started": _local_iso(self.started),
                "finished": _local_iso(time.time()) if status != "running" else None,
                "facts": self.facts,
                "layers": [self.results[n] for n in self.layers]}

    # -- events ----------------------------------------------------------- #
    def on_event(self, ev):
        kind = ev.get("type")
        if kind == "progress":
            self.events.on_log(f"  {ev.get('message')}")
        elif kind == "runtime":
            self.facts.update({k: v for k, v in ev.items() if k not in ("type", "t")})
            self.events.on_log(_runtime_line(self.facts))
        elif kind == "catalog":
            self.facts["catalog"] = bool(ev.get("ok"))
            if ev.get("ok"):
                self.events.on_log(f"  Read the service's layer list ({ev.get('count')} "
                                   "layers and tables) — layers are matched by name.")
            else:
                self.events.on_log("  Couldn't read the service's layer list "
                                   f"({ev.get('error')}) — using the known layer ids.")
        elif kind == "layer_start":
            self.current = ev.get("name")
            self.layer_started = time.time()
            r = self.results.get(self.current)
            if r is not None:
                r.update(status="exporting", started=_local_iso(ev.get("t") or time.time()))
            self.events.on_log(f"Exporting {self.current} "
                               f"({ev.get('index')} of {ev.get('total')})…")
            self.progress()
        elif kind == "layer_count":
            r = self.results.get(ev.get("name"))
            if r is not None:
                r["server_rows"] = ev.get("rows")
            self.events.on_log(f"  {ev.get('name')}: {int(ev.get('rows') or 0):,} rows on the server")
        elif kind == "layer_warning":
            self.events.on_log(f"  note: {ev.get('name')}: {ev.get('message')}")
        elif kind == "alive":
            self.progress()
        elif kind == "layer_done":
            self._place(ev)
        elif kind == "layer_failed":
            self._fail(ev.get("name"), ev.get("message") or "the export failed",
                       ev.get("detail"))
        elif kind == "fatal":
            self.events.on_log(f"ArcGIS layer export stopped: {ev.get('message')}")
            if ev.get("detail"):
                log.info("arcgis worker fatal detail:\n%s", ev.get("detail"))

    def _fail(self, name, message, detail=None):
        r = self.results.get(name)
        if r is not None:
            r.update(status="failed", message=message)
        if name == self.current:
            self.current, self.layer_started = None, None
            self.done_count += 1
        self.events.on_log(f"  {name}: FAILED — {message}")
        if detail:
            log.info("arcgis layer %s failed:\n%s", name, detail)
        self.record()
        self.progress()

    def _place(self, ev):
        name = ev.get("name")
        r = self.results.get(name)
        if r is None:
            return
        try:
            self._place_checked(ev, name, r)
        except ValueError as e:
            self._fail(name, str(e))
        except Exception as e:         # noqa: BLE001 — one layer's failure never ends the run
            log.exception("arcgis refresh: placing %s failed", name)
            hint = (" — close it in Excel and refresh this layer again"
                    if isinstance(e, PermissionError) else "")
            self._fail(name, f"couldn't add the export to the library "
                             f"({type(e).__name__}: {e}){hint}")

    def _backup_once(self):
        """Clear the previous generation and keep the current manifest, the
        first time this run replaces a file."""
        backup = backup_dir()
        if not self.backed_up:
            backup.mkdir(parents=True, exist_ok=True)
            _clear_files(backup)
            index = self.lib / crl.INDEX_NAME
            if index.is_file():
                shutil.copy2(index, backup / crl.INDEX_NAME)
            self.backed_up = True
        return backup

    def _place_checked(self, ev, name, r):
        staged = Path(ev.get("file") or "")
        info = inspect_export(staged, server_rows=ev.get("rows"))
        previous = self.inventory["present"].get(name)
        if previous is not None and not previous.is_file():
            previous = None
        old_header = None
        if previous is not None:
            try:
                old_header = read_header(previous)
            except Exception:          # silent-ok: the column diff is informational
                old_header = None
        target = self.lib / _target_name(name, self.inventory)
        kept = None
        if previous is not None:
            kept = self._backup_once() / previous.name
            os.replace(previous, kept)
        try:
            os.replace(staged, target)
        except OSError:
            if kept is not None:
                os.replace(kept, previous)
            raise
        exported_at = _local_iso(ev.get("started") or time.time())
        version = self.facts.get("arcgis_version")
        old_row = self.index.get(name)
        self.index[name] = {
            "file": target.name, "rows": info["data_rows"], "fields": info["fields"],
            "path": name, "source": ev.get("url") or "",
            "exported_at": exported_at,
            "exported_by": (f"TSMIS Exporter {_app_version()}"
                            + (f" · ArcGIS Pro {version}" if version else ""))}
        try:
            write_index(self.lib, self.index)
        except Exception:
            # The manifest must describe the files beside it: put the layer
            # back the way it was before reporting the failure.
            if old_row is None:
                self.index.pop(name, None)
            else:
                self.index[name] = old_row
            try:
                os.replace(target, staged)
                if kept is not None:
                    os.replace(kept, previous)
            except OSError as e:
                log.error("arcgis refresh: could not roll back %s (%s)", name, e)
            raise
        self.inventory["present"][name] = target
        changed = ""
        if old_header is not None and old_header != info["header"]:
            added = [c for c in info["header"] if c not in old_header]
            gone = [c for c in old_header if c not in info["header"]]
            changed = "columns changed" + (f" (+{', '.join(added)})" if added else "") \
                + (f" (−{', '.join(gone)})" if gone else "")
        r.update(status="exported", rows=info["data_rows"], fields=info["fields"],
                 size=target.stat().st_size, seconds=ev.get("seconds"),
                 exported_at=exported_at, file=target.name,
                 message=changed or None)
        extra = (f" — {changed}" if changed else "")
        rows_note = ""
        server = ev.get("rows")
        if server is not None and info["data_rows"] != server:
            rows_note = f" (the server counted {server:,})"
        self.events.on_log(f"  {name}: {info['data_rows']:,} rows{rows_note}, "
                           f"{info['fields']} columns in {ev.get('seconds')} s{extra}")
        self.current, self.layer_started = None, None
        self.done_count += 1
        self.record()
        self.progress()


def _runtime_line(facts):
    bits = []
    if facts.get("arcgis_version"):
        bits.append(f"ArcGIS Pro {facts['arcgis_version']}")
    if facts.get("license"):
        bits.append(f"license {facts['license']}")
    if facts.get("portal"):
        bits.append(("signed in to " if facts.get("signed_in") else "NOT signed in to ")
                    + str(facts["portal"]))
    elif facts.get("signed_in") is False:
        bits.append("NOT signed in")
    return "  " + (" · ".join(bits) if bits else "ArcGIS Pro started")


def resolve_layers(layers):
    """The layers to refresh, in library order: all 40 when `layers` is empty.
    Raises ValueError for a name that isn't a library layer."""
    if not layers:
        return [name for name, _id in LIBRARY_LAYERS]
    wanted = []
    for name in layers:
        if name not in LAYER_IDS:
            raise ValueError(f"Unknown layer: {name!r}")
        if name not in wanted:
            wanted.append(name)
    return [name for name, _id in LIBRARY_LAYERS if name in wanted]


def no_python_message():
    return ("ArcGIS Pro was not found on this PC. The layer refresh runs in "
            "ArcGIS Pro's own Python — on a PC with ArcGIS Pro installed, use "
            "“Choose python.exe…” if it isn't detected.")


def signin_hint(facts):
    """The one-line fix when ArcGIS Pro's sign-in is the likely cause: it is
    not signed in, or its active portal is not the TSMIS service's portal.
    "" when the sign-in looks right or is unknown."""
    portal = str(facts.get("portal") or "")
    service_host = SERVICE_DEFAULT.split("/")[2].lower()
    if facts.get("signed_in") is False:
        return ("ArcGIS Pro is not signed in" + (f" to {portal}" if portal else "")
                + ". Open ArcGIS Pro, sign in (tick “Sign me in automatically”), "
                  "close it, and try again.")
    if portal and service_host not in portal.lower():
        return (f"ArcGIS Pro's active portal is {portal}, not the TSMIS portal "
                f"(https://{service_host}/portal). Make that portal active in "
                "ArcGIS Pro (Settings ▸ Portals), sign in, and try again.")
    return ""


def refresh(layers=None, events=None, cancel_event=None, on_progress=None,
            python=None, launcher=None, service=SERVICE_DEFAULT):
    """Export `layers` (default: all 40) from the feature service with ArcGIS
    Pro and swap each finished one into the library. Returns a RefreshResult.
    Never raises for an ordinary failure (no ArcGIS Pro, sign-in, a bad
    layer) -- it says so in the result and the per-layer outcomes."""
    events = events or Events()
    try:
        names = resolve_layers(layers)
    except ValueError as e:
        return RefreshResult("error", str(e))
    if python is None:
        python, _how = arcgis_pro.resolve_python()
    if not python:
        return RefreshResult("error", no_python_message())
    lib = library_root()
    try:
        lib.mkdir(parents=True, exist_ok=True)
        staging_dir().mkdir(parents=True, exist_ok=True)
        _clear_files(staging_dir())
    except OSError as e:
        return RefreshResult("error", f"The layer library folder isn't writable ({e}).")
    run = _Refresher(names, events, on_progress, service)
    run.record()
    events.on_log("=" * 60)
    events.on_log(f"Refreshing {len(names)} ArcGIS layer(s) from {service}")
    events.on_log(f"ArcGIS Pro's Python: {python}")
    events.on_log("=" * 60)
    run.progress("Starting ArcGIS Pro…")
    request = {"protocol": 1, "mode": "export", "service": service,
               "layers": [{"name": n, "id": LAYER_IDS[n],
                           "out": str(staging_dir() / _target_name(n, run.inventory))}
                          for n in names]}
    try:
        outcome = arcgis_pro.run_worker(session_dir(), request, run.on_event,
                                        cancel_event, python, launcher=launcher)
    except Exception as e:                       # noqa: BLE001 — a launch failure is a result, not a crash
        log.exception("arcgis refresh: the worker could not run")
        outcome = SessionOutcome(fatal=f"ArcGIS Pro's Python could not be started "
                                       f"({type(e).__name__}: {e}).")
    _clear_files(staging_dir())
    return _finish(run, outcome)


def _finish(run, outcome):
    stopped = ("cancelled" if outcome.cancelled else
               f"stopped: {outcome.timed_out}" if outcome.timed_out else None)
    for name in run.layers:
        r = run.results[name]
        if r["status"] in ("queued", "exporting"):
            if stopped:
                r.update(status="cancelled" if outcome.cancelled else "failed",
                         message=stopped if not outcome.cancelled else None)
            else:
                r.update(status="failed", message=outcome.fatal or "not exported")
    exported = [n for n in run.layers if run.results[n]["status"] == "exported"]
    failed = [n for n in run.layers if run.results[n]["status"] == "failed"]
    if outcome.cancelled:
        status = "cancelled"
        message = (f"Layer refresh cancelled — {len(exported)} of {len(run.layers)} "
                   "layer(s) refreshed before it stopped.")
    elif not exported:
        status = "failed"
        why = outcome.fatal or outcome.timed_out or (
            run.results[failed[0]].get("message") if failed else "")
        message = "No layer was refreshed" + (f": {why}" if why else ".")
        if outcome.output_tail and not outcome.fatal:
            message += f"\n\nArcGIS Pro's Python said:\n{outcome.output_tail[-600:]}"
    elif failed:
        status = "partial"
        message = (f"Refreshed {len(exported)} of {len(run.layers)} layer(s); "
                   f"{len(failed)} failed: {', '.join(failed)}. Refresh them again "
                   "to retry.")
    else:
        status = "ok"
        message = f"Refreshed {len(exported)} layer(s)."
    if outcome.timed_out and status != "failed":
        message += f" The export stopped early ({outcome.timed_out})."
    # Only when nothing exported: a sign-in problem fails every layer, and a
    # partial run that read some layers was evidently signed in enough.
    hint = signin_hint(run.facts) if failed and not exported else ""
    if hint:
        message += f"\n\n{hint}"
    summary = run.summary(status, message)
    summary["output_tail"] = outcome.output_tail[-1500:] if status != "ok" else ""
    _atomic_json(work_dir() / _LAST_RUN, summary)
    run.events.on_log(message)
    return RefreshResult(status, message, summary["layers"], run.facts)


# --------------------------------------------------------------------------- #
# the ArcGIS Pro check
# --------------------------------------------------------------------------- #
def probe(events=None, cancel_event=None, python=None, launcher=None,
          service=SERVICE_DEFAULT):
    """Start ArcGIS Pro's Python, read the 27-row SHS Tolls layer and export it
    to a scratch file, then compare that file with the library's own copy.
    Returns (and records in last_probe.json) what it found. Never raises."""
    events = events or Events()
    how = "given"
    if python is None:
        python, how = arcgis_pro.resolve_python()
    result = {"when": _now_local().isoformat(), "ok": False, "python": python,
              "how": how, "service": service}
    if not python:
        result["message"] = no_python_message()
        _atomic_json(work_dir() / _LAST_PROBE, result)
        return result
    facts, layer = {}, {}

    def on_event(ev):
        kind = ev.get("type")
        if kind == "runtime":
            facts.update({k: v for k, v in ev.items() if k not in ("type", "t")})
            events.on_log(_runtime_line(facts))
        elif kind == "catalog":
            facts["catalog"] = bool(ev.get("ok"))
            facts["catalog_error"] = ev.get("error")
        elif kind == "progress":
            events.on_log(f"  {ev.get('message')}")
        elif kind in ("layer_count", "layer_done", "layer_failed"):
            layer.update({k: v for k, v in ev.items() if k not in ("type", "t")})
            layer["status"] = kind
        elif kind == "fatal":
            result["fatal"] = ev.get("message")

    events.on_log(f"Checking ArcGIS Pro ({python})…")
    scratch = session_dir().parent / "probe"
    try:
        scratch.mkdir(parents=True, exist_ok=True)
        _clear_files(scratch)
    except OSError as e:
        result["message"] = f"The layer library folder isn't writable ({e})."
        _atomic_json(work_dir() / _LAST_PROBE, result)
        return result
    out = scratch / f"{PROBE_LAYER}.xlsx"
    request = {"protocol": 1, "mode": "probe", "service": service,
               "probe_layer": {"name": PROBE_LAYER, "id": LAYER_IDS[PROBE_LAYER],
                               "out": str(out)}}
    try:
        outcome = arcgis_pro.run_worker(session_dir(), request, on_event, cancel_event,
                                        python, launcher=launcher,
                                        layer_limit=PROBE_LAYER_LIMIT_S)
    except Exception as e:                       # noqa: BLE001 — reported, never raised
        log.exception("arcgis probe: the worker could not run")
        outcome = SessionOutcome(fatal=f"ArcGIS Pro's Python could not be started "
                                       f"({type(e).__name__}: {e}).")
    result.update(facts=facts, layer=layer, cancelled=outcome.cancelled)
    if layer.get("status") == "layer_done":
        try:
            info = inspect_export(out, server_rows=layer.get("rows"))
            result["rows"] = info["data_rows"]
            result["fields"] = info["fields"]
            result.update(_compare_with_library(info["header"], out))
            result["ok"] = True
        except ValueError as e:   # silent-ok: becomes the check's own message, shown on the card and logged below
            result["message"] = f"{PROBE_LAYER} exported but the file is not usable: {e}"
        except Exception as e:                   # noqa: BLE001 — reported, never raised
            log.exception("arcgis probe: checking the export failed")
            result["message"] = (f"{PROBE_LAYER} exported but checking it failed "
                                 f"({type(e).__name__}: {e})")
    if not result.get("ok") and not result.get("message"):
        why = (layer.get("message") if layer.get("status") == "layer_failed" else None) \
            or result.get("fatal") or outcome.fatal or outcome.timed_out \
            or ("cancelled" if outcome.cancelled else "") or "the check did not finish"
        result["message"] = f"ArcGIS Pro could not read {PROBE_LAYER}: {why}"
        hint = signin_hint(facts)
        if hint:
            result["message"] += f"\n\n{hint}"
        if outcome.output_tail and not (layer or result.get("fatal")):
            result["message"] += f"\n\nArcGIS Pro's Python said:\n{outcome.output_tail[-600:]}"
    if result.get("ok"):
        result["message"] = (f"ArcGIS Pro read and exported {PROBE_LAYER} "
                             f"({result['rows']:,} rows, {result['fields']} columns) "
                             f"in {layer.get('seconds')} s.")
    _clear_files(scratch)
    _atomic_json(work_dir() / _LAST_PROBE, result)
    events.on_log(result["message"])
    return result


def _compare_with_library(header, probe_file):
    """How the probe's export lines up with the library's own copy of the same
    layer: the same columns, and the same value dialect (domain labels like
    'District 12' rather than bare codes)."""
    inv = crl.inventory()
    mine = inv["present"].get(PROBE_LAYER)
    out = {"dialect": _dialect(probe_file)}
    if mine is None:
        out["library_match"] = None
        return out
    try:
        theirs = read_header(mine)
    except Exception:          # silent-ok: informational comparison
        out["library_match"] = None
        return out
    out["library_match"] = theirs == header
    if theirs != header:
        out["library_diff"] = {"added": [c for c in header if c not in theirs],
                               "missing": [c for c in theirs if c not in header]}
    return out


def _dialect(path):
    """'labels' when the District column reads like 'District 12', 'codes'
    when it reads like '12', None when it can't tell."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:          # silent-ok: informational
        return None
    try:
        it = wb.worksheets[0].iter_rows(values_only=True)
        header = [str(c) if c is not None else "" for c in next(it, ())]
        if "District" not in header:
            return None
        i = header.index("District")
        for row in it:
            v = row[i] if i < len(row) else None
            if v is None or str(v).strip() == "":
                continue
            return "labels" if str(v).strip().upper().startswith("DISTRICT") else "codes"
        return None
    finally:
        wb.close()


# --------------------------------------------------------------------------- #
# what the Layers tab shows
# --------------------------------------------------------------------------- #
def library_status():
    """Every library layer with its file, manifest facts, size, the builds that
    read it and its last refresh outcome, plus the run records. Pure
    filesystem reads; never raises for an ordinary state."""
    lib = library_root()
    inv = crl.inventory(lib)
    index, created = read_index_full(lib)
    try:
        users = layer_users()
    except Exception as e:        # silent-ok: the column is informational
        log.info("arcgis_refresh: layer users unavailable (%s: %s)", type(e).__name__, e)
        users = {}
    run = last_run() or {}
    outcomes = {r.get("name"): r for r in run.get("layers") or [] if isinstance(r, dict)}
    rows = []
    total_size = 0
    for name, _id in LIBRARY_LAYERS:
        path = inv["present"].get(name)
        entry = index.get(name) or {}
        size = None
        if path is not None:
            try:
                size = path.stat().st_size
                total_size += size
            except OSError:      # silent-ok: vanished mid-listing
                size = None
        exported_at = entry.get("exported_at") or (created.isoformat() if created and entry else None)
        rows.append({
            "name": name, "file": path.name if path is not None else None,
            "present": path is not None, "rows": entry.get("rows"),
            "fields": entry.get("fields"), "size": size,
            "exported_at": str(exported_at) if exported_at else None,
            "exported_by": entry.get("exported_by"),
            "source": entry.get("source") or None,
            "used_by": users.get(name, []),
            "last": outcomes.get(name)})
    return {"root": str(lib), "service": SERVICE_DEFAULT, "layers": rows,
            "present": len(inv["present"]), "expected": len(LIBRARY_LAYERS),
            "missing": inv["missing"], "unknown": inv["unknown"],
            "index_present": inv["index"] is not None, "size": total_size,
            "backup": backup_dir().is_dir() and any(backup_dir().glob("*.xlsx")),
            "last_run": {k: v for k, v in run.items() if k != "layers"} or None,
            "last_probe": last_probe()}
