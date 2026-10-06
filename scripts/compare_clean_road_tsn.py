"""The SITE's Clean Road exports vs the TSN clean-road extracts (v0.48.0).

Nine "files"-kind comparisons over the regression-locked compare_core engine —
three per report (Highway / Intersection / Ramp):

  * <KIND>_VS_TSN        — the consolidated Excel export vs the TSN extract;
  * <KIND>_PDF_VS_TSN    — the consolidated PRINT vs the TSN extract;
  * <KIND>_PDF_VS_EXCEL  — the print vs the Excel export of the same run (the
                           site checked against itself).

Side A is always a workbook from `clean_road_consolidate` (sheet = the export's
own sheet, "Route" + the export's exact header). The TSN side is the raw
`CA HIGHWAYS` / `CA INTERSECTIONS` / `CA RAMPS` extract (`Sheet 1`) or the TSN
library's normalized copy (marker-gated). Role gates keep the sides honest: the
print flavors REQUIRE the PDF-conversion marker on side A, the Excel flavors
and the self-check's Excel side REFUSE it, and no side may be the app's own
ArcGIS build (`ArcGIS Build` marker) — that is a different comparison.

Row identity — the physical location both systems carry natively:
    route (name + suffix) · county · PM prefix · postmile (decimal-canonical)
    · roadbed (the R/L/X postmile suffix)
For Highway that is the segment's BEGIN postmile (the END is a compared field,
so a stretch the systems cut differently pairs instead of splitting); for
Intersection the intersection's postmile; for Ramp the ramp's.

Format-only reconciliations, all measured on the 2026-10-02 statewide pull
against the 2025-09 extracts and named on the Notes sheet:
  * dates compare as ISO dates (the site prints MM/DD/YYYY, Intersection dates
    pass the layer's own YY-MM-DD through; TSN stores dates);
  * amounts compare as canonical numbers ('02' = '2', '9.60' = '9.6');
  * county codes drop the TASAS trailing period the site writes for the
    two-letter counties ('SD.' = 'SD');
  * Highway: the site writes the equation marker 'E' into the postmile-suffix
    column on equation rows (THY_EQUATE_CODE = E on every one), so it is not a
    roadbed for the KEY — the column itself is still compared; a blank
    THY_TOLL_FOREST_CODE is TSN's code 0 (neither toll nor forest);
  * Intersection: TSN prefixes its design codes with the domain letter
    ('IT' = the site's 'T'); the control type applies the TSNR/MIRE crosswalk
    the Intersection Detail comparison already uses (TSN's signal sub-types
    J–P fold to the Signalized code S on both sides);
  * Ramp: TSN prefixes its design codes ('RD' = 'D'); the site's On/Off 'Z'
    is TSN's 'OTH'.
Runs of spaces compare as one space in every flavor (the engine's Excel-TRIM
rule) — which is what lets the print, whose HTML collapses them, compare
against the Excel file and the TSN extract, which keep them; the print flavors
also fold tabs and line breaks the same way.

CONTEXT columns (shown with both sides' values, never counted): the columns
the site has no source for (blank in its file by design), TSN's bookkeeping
(ids, lifecycle, create/update), the TASAS change-tracking columns, the
extract date, Highway's two offset columns (each system's own cumulative), the
Intersection connection id (a different id space) and the Ramp design
description (each system's own wording of the compared design code).

TSN is FROZEN at the September-2025 cutover (roadmap D5): every difference is
TSMIS-vs-the-2025-snapshot drift, never by itself a defect. Console-free.
"""
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

try:
    from openpyxl import load_workbook
    _DEPS_OK = True
except ImportError:
    _DEPS_OK = False

import clean_highway_columns as chc
import clean_road_columns as crc
import compare_clean_highway_tsn as cht
import compare_tsn_common as ctc
from compare_core import CompareSchema
from compare_intersection_detail_tsn import _norm_control_type
from paths import today_str
from pdf_table_lib import pdf_source_marker_state

# CMP-AUD-037: the Intersection / Ramp library slots are stamped with this
# version (tsn_load_clean_road mirrors it; the catalog's normalization_version
# mirrors both). Highway shares the ArcGIS comparison's slot and marker.
NORMALIZATION_VERSION = 1
_HIGHWAY_TSN_MARKER = cht.NORMALIZATION_VERSION

_WS = re.compile(r"\s+")


def _s(v):
    return "" if v is None else str(v).strip()


def _date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return ctc.iso_date(v)


def _county(v):
    return _s(v).upper().rstrip(".")


@dataclass(frozen=True)
class _Profile:
    """One report's comparison rules (see the module docstring)."""
    spec: object
    context: tuple
    dates: frozenset
    amounts: frozenset
    text: tuple                 # free-text columns (representation disclosure)
    tsn_design_prefix: str      # TSN's domain letter on the design code, or ""
    design_col: str
    notes: tuple


def _is_numeric_col(name):
    return name.endswith(("_AMT", "_ID", "_ADT")) or name in (
        "THY_TOLL_FOREST_CODE", "THY_CURB_LANDSCAPE_CODE", "THY_MAINT_SVC_LVL_CODE",
        "THY_NATIONAL_LANDS_CODE", "THY_SCENIC_FREEWAY_CODE", "RAM_ADT",
        "INX_MAINLINE_ADT", "INX_XSTREET_ADT")


def _profile(spec, *, extra_context, design_col, tsn_design_prefix, text, notes):
    header = spec.header
    dates = frozenset(c for c in header if c.endswith("_DATE"))
    amounts = frozenset(c for c in header if _is_numeric_col(c) and c not in dates)
    context = tuple(c for c in header
                    if c in set(spec.unsourced) | set(extra_context)
                    or c in spec.site_only or c.endswith("_EXTRACT_DATE"))
    return _Profile(spec=spec, context=context, dates=dates, amounts=amounts,
                    text=text, tsn_design_prefix=tsn_design_prefix,
                    design_col=design_col, notes=notes)


_COMMON_NOTES = (
    "The TSN side is FROZEN at the September-2025 cutover (TSMIS replaced "
    "TSN), so every difference here is TSMIS-vs-that-snapshot drift — read the "
    "counts as change since then, never by themselves as defects.",
    "Dates compare as ISO dates (format never counts); amounts as canonical "
    "numbers ('02' = '2', '9.60' = '9.6'); county codes drop the trailing "
    "period the site writes for two-letter counties ('SD.' = 'SD').",
)

_HIGHWAY = _profile(
    crc.HIGHWAY,
    extra_context=tuple(c for c in chc.CONTEXT_COLUMNS if c in chc.HEADER),
    design_col="", tsn_design_prefix="",
    text=("THY_LANDMARK_SHORT_DESC",),
    notes=_COMMON_NOTES + (
        "Rows pair on route + county + PM prefix + BEGIN postmile + roadbed "
        "(the R/L/X postmile suffix). The END postmile and length are compared "
        "fields, so a stretch the two systems cut differently pairs on its "
        "begin and shows its end as a difference instead of splitting into two "
        "one-sided rows. The site segments the road at the Highway Log's event "
        "points, TSN at its own record breaks, so one-sided rows are mostly "
        "segmentation; TSN also carries the unconstructed routes the site's "
        "file doesn't.",
        "The site writes the equation marker 'E' into THY_PM_SUFFIX_CODE on "
        "equation rows (every one also carries THY_EQUATE_CODE = E); it is not a "
        "roadbed, so the row key treats it as none. The column is still "
        "compared, so the marker shows against TSN's blank.",
        "THY_TOLL_FOREST_CODE: the site leaves the cell blank where no toll or "
        "forest span exists; TSN writes its code 0 for that. A blank compares as "
        "0 on both sides.",
        "CONTEXT (shown, never counted): the 22 columns the site has no source "
        "for (blank in its file by design — ids, lifecycle dates, the four "
        "change-tracking flags, the ADT profile trio, maintenance service, the "
        "federal-aid trio, national lands, scenic freeway), TSN's bookkeeping, "
        "THY_LAST_SIG_CHG_DATE and THY_FUNCTIONAL_CLASS_CODE (TSN's own "
        "change-tracking date; TSN leaves the class blank), the two offset "
        "columns (each system's own cumulative) and the extract date.",
    ))

_INTERSECTION = _profile(
    crc.INTERSECTION,
    extra_context=("INX_CONNECTION_ID",),
    design_col="INX_DESIGN_CODE", tsn_design_prefix="I",
    text=("INX_INTERSECTION_NAME",),
    notes=_COMMON_NOTES + (
        "Rows pair on route + county + PM prefix + the intersection's postmile "
        "+ roadbed (the R/L postmile suffix).",
        "INX_DESIGN_CODE: TSN prefixes its intersection design codes with the "
        "domain letter 'I' ('IT' = the site's 'T'); the letter is dropped from "
        "the TSN value before comparing.",
        "INX_CONTROL_CODE applies the TSNR/MIRE crosswalk the Intersection "
        "Detail comparison uses: TSN's signal sub-types J, K, L, M, N and P and "
        "TSMIS's S all fold to S (Signalized) on both sides.",
        "The site passes the intersection layer's own YY-MM-DD dates through "
        "('73-10-19'); they read as 1973-10-19 (2-digit years below 30 are 20xx).",
        "CONTEXT (shown, never counted): the 10 columns the site has no source "
        "for (placement ids, create/update, end date, both ADTs, the LSC date), "
        "INX_CONNECTION_ID (the site numbers intersections in a different id "
        "space) and the extract date.",
    ))

_RAMP = _profile(
    crc.RAMP,
    extra_context=("RAM_DESIGN_DESC",),
    design_col="RAM_DESIGN_CODE", tsn_design_prefix="R",
    text=("RAM_DESCRIPTION",),
    notes=_COMMON_NOTES + (
        "Rows pair on route + county + PM prefix + the ramp's postmile + "
        "roadbed (the R/L postmile suffix) — unique on both sides statewide.",
        "RAM_DESIGN_CODE: TSN prefixes its ramp design codes with the domain "
        "letter 'R' ('RD' = the site's 'D'); the letter is dropped from the TSN "
        "value before comparing. RAM_DESIGN_DESC is each system's own wording "
        "of that code (the site's ArcGIS domain vs TSN's code table) and is "
        "shown as context.",
        "RAM_ON_OFF_CODE: the site writes 'Z' for the 'other' class, TSN 'OTH'; "
        "they compare equal.",
        "RAM_BEGIN_DATE: the site fills it from the ramp layer's inventory start "
        "date, which is the date TSN keeps in RAM_CHANGE_DATE (TSN's "
        "RAM_BEGIN_DATE is its own record-version date). The column is compared "
        "as named, so this mapping difference shows on most ramps.",
        "CONTEXT (shown, never counted): the two site-only columns (RAM_HPMS_ID, "
        "RAM_RAMP_ROUTE_NAME — TSN has neither), the 13 columns the site has no "
        "source for (ids, offset, create/update, ADT, change date, connection "
        "id), the design description and the extract date.",
    ))

_PROFILES = {p.spec.key: p for p in (_HIGHWAY, _INTERSECTION, _RAMP)}


# --------------------------------------------------------------------------- #
# row projection
# --------------------------------------------------------------------------- #
def _value(profile, name, raw, side, collapse):
    """One cell as compared. `side` is "site" or "tsn"; `collapse` folds
    whitespace runs (the print flavors)."""
    if name in profile.dates:
        return _date(raw)
    if name.endswith("_COUNTY_CODE"):
        return _county(raw)
    if name == "THY_TOLL_FOREST_CODE":
        return cht._norm_amount(raw) or "0"
    if name == "THY_CHANGE_PER_MILE_AMT":
        return cht._norm_cell(name, raw)
    if name in profile.amounts:
        return cht._norm_amount(raw)
    # The installed-Excel reading (CMP-AUD-197): the export serializes control
    # characters as OOXML `_xHHHH_` escapes; decoded, any such character stays
    # compared content (the Ramp Detail vs-TSN rule).
    text = ctc.decode_ooxml_escapes(_s(raw))
    if collapse:
        text = _WS.sub(" ", text)
    if name == profile.design_col and side == "tsn" and profile.tsn_design_prefix:
        if len(text) == 2 and text[0].upper() == profile.tsn_design_prefix:
            text = text[1:]
    elif name == "INX_CONTROL_CODE":
        text = _norm_control_type(text)
    elif name == "RAM_ON_OFF_CODE" and text.upper() == "Z":
        text = "OTH"
    return text


def _self_value(profile, name, raw):
    """The PDF-vs-Excel projection: both sides are the SITE's own renders, so
    every value is verbatim except what one render structurally cannot carry —
    the whitespace runs the HTML print collapses, the Excel export's OOXML
    escapes and edge tab padding (the CMP-AUD-197 same-source rule) — and a
    date the workbook may have typed."""
    if isinstance(raw, (datetime, date)):
        return _date(raw)
    return _WS.sub(" ", ctc.same_source_render_text(_s(raw)))


_ROUTE_RE = re.compile(r"(\d+)([A-Z]*)")


def route_identity(name, suffix):
    """One route token from a row's route-name + route-suffix cells, both sides
    alike. TSN writes the bare number ('005') with the suffix in its own column;
    the site's Clean Road Highway file writes the selected route LABEL ('005S')
    as the name and fills the suffix column on some rows only — so a suffix
    already in the name is never appended twice ('005S' + 'S' = '005S')."""
    text = _s(name).upper()
    if text.endswith(".0"):
        text = text[:-2]
    sfx = _s(suffix).upper()
    m = _ROUTE_RE.fullmatch(text)
    if not m:
        return text + sfx
    digits, own = m.groups()
    base = f"{int(digits):03d}"
    if own:
        return base + own + (sfx if sfx and sfx != own else "")
    return base + sfx


def _key(profile, vals, source_hint, route=None):
    """(route, the physical key) for one row. `route` overrides the row's own
    route cells (the cross-environment path keys on the file's route token,
    which the engine requires the identity to equal)."""
    spec = profile.spec
    g = dict(zip(spec.header, vals))
    if route is None:
        route = route_identity(g.get(spec.route_col), g.get(spec.suffix_col))
    roadbed = _s(g.get(spec.pm_suffix_col)).upper()
    if roadbed not in ("R", "L", "X"):
        roadbed = ""                     # the site's equation marker 'E' is no roadbed
    return route, cht._physical_span_key(
        route, _county(g.get(spec.county_col)), _s(g.get(spec.pm_prefix_col)).upper(),
        g.get(spec.pm_col), roadbed, source_hint)


def _row(profile, vals, side, source_hint, *, same_source=False, collapse=False):
    """[route, *header values] with the postmile cell as the physical key."""
    spec = profile.spec
    route, key = _key(profile, vals, source_hint)
    out = [route]
    for name, raw in zip(spec.header, vals):
        if name == spec.pm_col:
            out.append(key)
        elif same_source:
            out.append(_self_value(profile, name, raw))
        else:
            out.append(_value(profile, name, raw, side, collapse))
    return out


# --------------------------------------------------------------------------- #
# loaders + role gates
# --------------------------------------------------------------------------- #
def _open(path):
    name = Path(path).name
    try:
        return name, load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Could not open {name}: {type(e).__name__}: {e}")


def load_site(path, profile, *, pdf, same_source=False, collapse=False):
    """Side A (and the self-check's side B): a consolidated SITE export —
    `clean_road_consolidate`'s sheet with "Route" + the export's exact header.
    `pdf` True = the print consolidation (marker required), False = the Excel
    one (marker refused)."""
    spec = profile.spec
    label = f"{spec.label} ({'PDF' if pdf else 'Excel'})"
    if pdf:
        ctc.require_pdf_source(path, f"TSMIS {label}", spec.label)
    else:
        ctc.reject_pdf_source(path, f"TSMIS {label}", spec.label)
    name, wb = _open(path)
    try:
        if chc.ARC_MARKER_SHEET in wb.sheetnames:
            raise ValueError(
                f"{name} is the app's own ArcGIS-built Clean Road workbook, not "
                "the site's export — pick the consolidated Clean Road export "
                "(the Consolidate tab, or the matrix builds it).")
        if spec.sheet not in wb.sheetnames:
            raise ValueError(f"{name} has no '{spec.sheet}' sheet — pick the "
                             f"consolidated {label} workbook.")
        it = wb[spec.sheet].iter_rows(values_only=True)
        header = [_s(c) for c in (next(it, None) or ())]
        while header and not header[-1]:
            header.pop()
        if header != ["Route"] + list(spec.header):
            raise ValueError(
                f"{name} is not a consolidated {label} workbook with the "
                f"{len(spec.header)}-column export header — re-consolidate it "
                "with this version.")
        rows = []
        for r in it:
            if not ctc.row_has_data(r):
                continue
            vals = list(r[1:1 + len(spec.header)])
            vals += [None] * (len(spec.header) - len(vals))
            rows.append(_row(profile, vals, "site", f"{name} ({spec.sheet})",
                             same_source=same_source, collapse=collapse))
        return rows
    finally:
        wb.close()


def _tsn_values(profile, raw):
    """A TSN row (in its own header order) onto the site header order: the
    site-only columns are simply absent on the TSN side."""
    spec = profile.spec
    by_name = dict(zip(spec.tsn_header, raw))
    return [by_name.get(c) for c in spec.header]


def load_tsn(path, profile, *, collapse=False):
    """Side B: the raw TSN extract (`Sheet 1`, exact header) or the TSN
    library's normalized copy (marker-gated). A PDF conversion, the ArcGIS build
    or a consolidated site export is refused — none of them is TSN."""
    spec = profile.spec
    if pdf_source_marker_state(path) != 0:
        raise ValueError(
            f"{Path(path).name} is one of this app's Clean Road PDF conversions, "
            f"so it cannot stand as the TSN side — pick the TSN {spec.tsn_label} "
            "extract or the TSN library's normalized copy.")
    name, wb = _open(path)
    try:
        if chc.ARC_MARKER_SHEET in wb.sheetnames:
            raise ValueError(
                f"{name} is an ArcGIS-built Clean Road workbook, so it cannot "
                f"stand as the TSN side — pick the TSN {spec.tsn_label} extract or "
                "the TSN library's normalized copy.")
        if spec.normalized_sheet in wb.sheetnames:
            it = wb[spec.normalized_sheet].iter_rows(values_only=True)
            header = [_s(c) for c in (next(it, None) or ())]
            ctc.require_shared_header_prefix(header, spec.tsn_header, (), name,
                                             spec.label)
            version = (_HIGHWAY_TSN_MARKER if spec is crc.HIGHWAY
                       else NORMALIZATION_VERSION)
            ctc.require_current_normalization(
                wb, name, version, "pre-v1: no in-workbook normalization marker")
            return [_row(profile, _tsn_values(profile, list(r)), "tsn",
                         f"{name} ({spec.normalized_sheet})", collapse=collapse)
                    for r in it if ctc.row_has_data(r)]
    finally:
        wb.close()
    return tsn_rows_from_raw(path, profile, collapse=collapse)


def required_identity(spec):
    """The TSN identity claims a raw row must carry."""
    return (spec.county_col, spec.route_col, spec.pm_col)


def tsn_rows_from_raw(path, profile, *, collapse=False):
    """Every row of the exact raw TSN statewide extract."""
    spec = profile.spec
    with ctc.exact_raw_rows(path, crc.TSN_RAW_SHEET, spec.tsn_header,
                            f"Clean Road {spec.tsn_label}",
                            required_nonblank=required_identity(spec)) as (_h, rows_in):
        return [_row(profile, _tsn_values(profile, list(r)), "tsn",
                     f"{Path(path).name} ({crc.TSN_RAW_SHEET})", collapse=collapse)
                for r in rows_in]


# --------------------------------------------------------------------------- #
# schemas + adapters
# --------------------------------------------------------------------------- #
def _schema(profile, side_a, side_b, *, notes_title, notes, context, representation):
    spec = profile.spec
    header = list(spec.header)
    widths = {c: 26 for c in profile.text}
    return CompareSchema(
        report_name=spec.label,
        header=header,
        side_a=side_a,
        side_b=side_b,
        id_noun=spec.noun,
        id_noun_plural=spec.noun_plural,
        pair_noun="postmile",
        sides_noun="systems" if side_b == "TSN" else "renders",
        date_fields=tuple(c for c in header if c in profile.dates),
        data_widths=widths,
        cmp_widths={c: 30 for c in profile.text},
        one_sided_note_extra=(f" ({spec.noun_plural} one side carries at a "
                              "physical location the other doesn't)"),
        key_field=header.index(spec.pm_col),
        context_fields=context,
        context_header_fill="808080" if context else "",
        legend_writer=ctc.make_notes_writer(notes_title, notes),
        representation_fields=representation,
    )


_SELF_NOTES = (
    "Both sides are the SITE's own renders of the same run: the print (read "
    "back from its PDF) and the Excel export. Every value compares verbatim; "
    "a run of spaces counts as one space (the HTML print collapses them, the "
    "Excel file keeps them).",
    "The print leaves out the columns the site has no source for; they are "
    "blank in the Excel file too, so they compare equal here.",
    "The Excel file writes a few control characters as escapes ('_x000d_' is "
    "a carriage return) and can pad a cell's edge with tabs; the print cannot "
    "carry either, so they are decoded away and never count (the owner's "
    "PDF-vs-Excel ruling, CMP-AUD-197).",
)


class _CleanRoadCompare:
    """One Clean Road file-vs-file comparison: compare(path_a, path_b, …) +
    suggest_name(path_a), with the side labels carried to the workbook and
    the Compare tab's file pickers."""

    def __init__(self, profile, flavor):
        spec = profile.spec
        self._profile = profile
        self._flavor = flavor
        if flavor == "excel_tsn":
            self.file_a_label, self.file_b_label = "TSMIS", "TSN"
            self.REPORT_NAME = f"{spec.label} — TSMIS vs TSN"
            tag, collapse = "TSMIS_vs_TSN", False
        elif flavor == "pdf_tsn":
            self.file_a_label, self.file_b_label = "TSMIS (PDF)", "TSN"
            self.REPORT_NAME = f"{spec.label} — TSMIS (PDF) vs TSN"
            tag, collapse = "TSMIS_PDF_vs_TSN", True
        elif flavor == "pdf_excel":
            self.file_a_label, self.file_b_label = "TSMIS (PDF)", "TSMIS (Excel)"
            self.REPORT_NAME = f"{spec.label} — TSMIS PDF vs Excel"
            tag, collapse = "TSMIS_PDF_vs_Excel", True
        else:
            raise ValueError(f"unknown Clean Road comparison flavor {flavor!r}")
        self._tag = f"{tag}_{spec.label.replace(' ', '')}"
        self._collapse = collapse
        if flavor == "pdf_excel":
            self._schema = _schema(
                profile, self.file_a_label, self.file_b_label,
                notes_title=f"{spec.label} — TSMIS PDF vs Excel: comparison notes",
                notes=_SELF_NOTES, context=(), representation=())
        else:
            lines = profile.notes + ((
                "This side was read back from the site's PRINT, whose HTML "
                "collapses runs of spaces; a run of spaces compares as one space "
                "in every flavor, and this flavor folds tabs and line breaks the "
                "same way on both sides.",) if collapse else (
                "The Excel export writes a few control characters as escapes "
                "('_x000d_' is a carriage return); they are read decoded, the way "
                "installed Excel shows the cell, and the character itself still "
                "compares.",))
            self._schema = _schema(
                profile, self.file_a_label, self.file_b_label,
                notes_title=f"{self.REPORT_NAME}: comparison notes",
                notes=lines, context=profile.context,
                representation=profile.text)

    def suggest_name(self, path_a=None):
        return f"{self._tag}_Comparison {today_str()}.xlsx"

    def _load_pair(self, path_a, path_b):
        p = self._profile
        if self._flavor == "pdf_excel":
            rows_a = load_site(path_a, p, pdf=True, same_source=True)
            rows_b = load_site(path_b, p, pdf=False, same_source=True)
        else:
            rows_a = load_site(path_a, p, pdf=(self._flavor == "pdf_tsn"),
                               collapse=self._collapse)
            rows_b = load_tsn(path_b, p, collapse=self._collapse)
        return rows_a, rows_b, None

    def compare(self, path_a, path_b, out_path, events=None, confirm_overwrite=None,
                mode="formulas", commit_guard=None):
        return ctc.run_files_compare(
            self._schema, path_a, path_b, out_path,
            banner=(f"{self._profile.spec.label} Comparison — {self.file_a_label} "
                    f"vs {self.file_b_label}"),
            has_route=True, loader=self._load_pair, deps_ok=_DEPS_OK,
            deps_msg="Required components are missing (openpyxl).",
            side_a=self.file_a_label, side_b=self.file_b_label,
            events=events, confirm_overwrite=confirm_overwrite, mode=mode,
            commit_guard=commit_guard)


HIGHWAY_VS_TSN = _CleanRoadCompare(_HIGHWAY, "excel_tsn")
HIGHWAY_PDF_VS_TSN = _CleanRoadCompare(_HIGHWAY, "pdf_tsn")
HIGHWAY_PDF_VS_EXCEL = _CleanRoadCompare(_HIGHWAY, "pdf_excel")
INTERSECTION_VS_TSN = _CleanRoadCompare(_INTERSECTION, "excel_tsn")
INTERSECTION_PDF_VS_TSN = _CleanRoadCompare(_INTERSECTION, "pdf_tsn")
INTERSECTION_PDF_VS_EXCEL = _CleanRoadCompare(_INTERSECTION, "pdf_excel")
RAMP_VS_TSN = _CleanRoadCompare(_RAMP, "excel_tsn")
RAMP_PDF_VS_TSN = _CleanRoadCompare(_RAMP, "pdf_tsn")
RAMP_PDF_VS_EXCEL = _CleanRoadCompare(_RAMP, "pdf_excel")


def profile_for(key):
    """The comparison profile for a Clean Road export key (Excel or print)."""
    spec = crc.BY_KEY.get(key) or crc.BY_PDF_KEY.get(key)
    return _PROFILES[spec.key] if spec is not None else None


__all__ = ["NORMALIZATION_VERSION", "HIGHWAY_VS_TSN", "HIGHWAY_PDF_VS_TSN",
           "HIGHWAY_PDF_VS_EXCEL", "INTERSECTION_VS_TSN", "INTERSECTION_PDF_VS_TSN",
           "INTERSECTION_PDF_VS_EXCEL", "RAMP_VS_TSN", "RAMP_PDF_VS_TSN",
           "RAMP_PDF_VS_EXCEL", "load_site", "load_tsn", "tsn_rows_from_raw",
           "profile_for"]
