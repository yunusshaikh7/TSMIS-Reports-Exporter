"""The second-edition comparisons of the three AGGREGATE summaries (v0.48.0).

Ramp Summary gained its Excel edition and Intersection / Highway Summary their
print editions in the export long ago; v0.48.0 consolidates each into the SAME
workbook as its sibling (`consolidate_ramp_summary_excel`,
`consolidate_tsmis_intersection_summary_pdf`,
`consolidate_tsmis_highway_summary_pdf`). Six "files"-kind comparisons ride on
that, all over the regression-locked compare_core engine:

  * vs TSN (one per family) — the family's own aggregate vs-TSN comparison
    (`compare_<family>_tsn`: its loaders, its schema, its familiar "Summary by
    Category" sheet and censused partition notes), with the second edition as
    the TSMIS side. Only the side label and the banner change, so a run
    against the same day reads exactly what the first edition reads.
  * PDF vs Excel (one per family) — the two TSMIS renders of the same run,
    route by route: one row per route, every category column compared
    verbatim. The PDF side must carry the PDF-conversion marker and the Excel
    side must not (CMP-AUD-066), so neither can stand in for the other.

Console-free.
"""
from dataclasses import replace
from pathlib import Path

try:
    from openpyxl import load_workbook
    _DEPS_OK = True
except ImportError:
    _DEPS_OK = False

import compare_highway_summary_tsn as _hs
import compare_intersection_summary_tsn as _is
import compare_ramp_summary_tsn as _rs
import compare_tsn_common as ctc
import consolidate_ramp_summary as _crs
import highway_summary_columns as hsc
import summary_layout
from compare_core import CompareSchema
from compare_env import HS_HEADER, IS_HEADER, RS_HEADER, _norm_route_key
from paths import today_str


# --------------------------------------------------------------------------- #
# vs TSN: the family's own comparison, the second edition as the TSMIS side
# --------------------------------------------------------------------------- #
class _SummaryVsTsn:
    """compare(tsmis_path, tsn_path, …) for one family's second edition."""

    def __init__(self, family, label, side_a, tag, *, pdf):
        self._family = family                 # "ramp" | "intersection" | "highway"
        self.REPORT_NAME = label
        self.file_a_label = side_a
        self.file_b_label = "TSN"
        self._tag = tag
        self._pdf = pdf

    def suggest_name(self, tsmis_path=None):
        return f"{self._tag}_Comparison {today_str()}.xlsx"

    def _gate(self, path):
        noun = self.REPORT_NAME.split(" — ")[0]
        if self._pdf:
            ctc.require_pdf_source(path, self.file_a_label, noun)
        else:
            ctc.reject_pdf_source(path, self.file_a_label, noun)

    def compare(self, tsmis_path, tsn_path, out_path, events=None,
                confirm_overwrite=None, mode="formulas", commit_guard=None):
        notes, footnotes = [], {}
        if self._family == "ramp":
            mod = _rs
            writer = summary_layout.make_extra_sheet_writer(
                _rs._SPEC, footnote_values=footnotes, extra_notes=notes)

            def pair(a, b):
                self._gate(a)
                return _rs._load_pair(a, b, footnote_sink=footnotes,
                                      note_sink=notes, events=events)
            measure = "statewide category counts"
        elif self._family == "intersection":
            mod = _is
            writer = summary_layout.make_extra_sheet_writer(_is._SPEC, extra_notes=notes)

            def pair(a, b):
                self._gate(a)
                return _is._load_pair(a, b, note_sink=notes, events=events)
            measure = "statewide category counts"
        else:
            mod = _hs
            writer = summary_layout.make_extra_sheet_writer(_hs._SPEC, extra_notes=notes)

            def pair(a, b):
                self._gate(a)
                return _hs._load_pair(a, b, note_sink=notes, events=events)
            measure = "statewide category miles"
        schema = replace(mod._SCHEMA, side_a=self.file_a_label,
                         extra_sheet_writer=writer)
        return ctc.run_files_compare(
            schema, tsmis_path, tsn_path, out_path,
            banner=(f"{mod.REPORT_NAME} Comparison — {self.file_a_label} vs TSN "
                    f"({measure})"),
            has_route=False, loader=pair, deps_ok=mod._DEPS_OK,
            deps_msg="Required components are missing (pdfplumber, openpyxl).",
            side_a=self.file_a_label, side_b="TSN",
            events=events, confirm_overwrite=confirm_overwrite, mode=mode,
            commit_guard=commit_guard)


RAMP_SUMMARY_EXCEL_VS_TSN = _SummaryVsTsn(
    "ramp", "Ramp Summary — TSMIS (Excel) vs TSN", "TSMIS (Excel)",
    "TSMIS_Excel_vs_TSN_RampSummary", pdf=False)
INTERSECTION_SUMMARY_PDF_VS_TSN = _SummaryVsTsn(
    "intersection", "Intersection Summary — TSMIS (PDF) vs TSN", "TSMIS (PDF)",
    "TSMIS_PDF_vs_TSN_IntersectionSummary", pdf=True)
HIGHWAY_SUMMARY_PDF_VS_TSN = _SummaryVsTsn(
    "highway", "Highway Summary — TSMIS (PDF) vs TSN", "TSMIS (PDF)",
    "TSMIS_PDF_vs_TSN_HighwaySummary", pdf=True)


# --------------------------------------------------------------------------- #
# PDF vs Excel: the two renders, route by route
# --------------------------------------------------------------------------- #
def _cell(v):
    if isinstance(v, str):
        return v.strip() or None
    return v


def _rows_ramp(path):
    """The consolidated Ramp Summary per-route sheet as [route, *RS fields]
    (the cross-environment row shape). Row 1 = group headers, row 2 = the
    column display names, data from row 3."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if _crs.SUMMARY_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"{Path(path).name} has no '{_crs.SUMMARY_SHEET_NAME}' "
                             "sheet — pick a consolidated Ramp Summary workbook.")
        it = wb[_crs.SUMMARY_SHEET_NAME].iter_rows(values_only=True)
        next(it, None)
        names = [str(h).strip() if h is not None else "" for h in (next(it, None) or ())]
        want = RS_HEADER
        missing = [h for h in want if h not in names]
        if missing:
            raise ValueError(f"{Path(path).name} isn't a consolidated Ramp Summary "
                             f"workbook (no {missing[0]!r} column) — consolidate first.")
        idx = [names.index(h) for h in want]
        rows = []
        for r in it:
            if not r or all(c is None for c in r):
                continue
            vals = [_cell(r[i]) if i < len(r) else None for i in idx]
            vals[0] = _norm_route_key(vals[0])
            rows.append(vals)
        return rows
    finally:
        wb.close()


def _rows_flat(path, sheet, header, noun):
    """A consolidated per-route summary sheet whose header IS the compared
    row shape (Intersection / Highway Summary): [route, *fields]."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"{Path(path).name} has no '{sheet}' sheet — pick a "
                             f"consolidated {noun} workbook.")
        it = wb[sheet].iter_rows(values_only=True)
        names = [str(h).strip() if h is not None else "" for h in (next(it, None) or ())]
        if names[:len(header)] != list(header):
            raise ValueError(f"{Path(path).name} isn't a consolidated {noun} workbook "
                             "with this version's columns — re-consolidate it.")
        rows = []
        for r in it:
            if not r or all(c is None for c in r):
                continue
            vals = [_cell(v) for v in list(r[:len(header)])]
            vals += [None] * (len(header) - len(vals))
            vals[0] = _norm_route_key(vals[0])
            rows.append(vals)
        return rows
    finally:
        wb.close()


_SELF_NOTES = (
    "Both sides are TSMIS renders of the same run: the print (read back from "
    "its PDF) and the Excel export, each consolidated into the same per-route "
    "workbook. One row per route; every category compares verbatim.",
)


class _SummaryPdfVsExcel:
    """compare(pdf_path, excel_path, …) for one family's two renders."""

    def __init__(self, label, noun, tag, header, loader):
        self.REPORT_NAME = label
        self.file_a_label = "TSMIS (PDF)"
        self.file_b_label = "TSMIS (Excel)"
        self._noun = noun
        self._tag = tag
        self._loader = loader
        self._schema = CompareSchema(
            report_name=label, header=list(header),
            side_a=self.file_a_label, side_b=self.file_b_label,
            id_noun="route", id_noun_plural="routes", sides_noun="renders",
            scope_flat="All routes (one row per route)",
            one_sided_note_extra=" (a route one render carries and the other doesn't)",
            legend_writer=ctc.make_notes_writer(f"{label}: comparison notes",
                                                _SELF_NOTES))

    def suggest_name(self, path_a=None):
        return f"{self._tag}_Comparison {today_str()}.xlsx"

    def _load_pair(self, path_a, path_b):
        ctc.require_pdf_source(path_a, self.file_a_label, self._noun)
        ctc.reject_pdf_source(path_b, self.file_b_label, self._noun)
        return self._loader(path_a), self._loader(path_b), None

    def compare(self, path_a, path_b, out_path, events=None, confirm_overwrite=None,
                mode="formulas", commit_guard=None):
        return ctc.run_files_compare(
            self._schema, path_a, path_b, out_path,
            banner=f"{self._noun} Comparison — TSMIS (PDF) vs TSMIS (Excel)",
            has_route=False, loader=self._load_pair, deps_ok=_DEPS_OK,
            deps_msg="Required components are missing (openpyxl).",
            side_a=self.file_a_label, side_b=self.file_b_label,
            events=events, confirm_overwrite=confirm_overwrite, mode=mode,
            commit_guard=commit_guard)


RAMP_SUMMARY_PDF_VS_EXCEL = _SummaryPdfVsExcel(
    "Ramp Summary — TSMIS PDF vs Excel", "Ramp Summary",
    "TSMIS_PDF_vs_Excel_RampSummary", RS_HEADER, _rows_ramp)
INTERSECTION_SUMMARY_PDF_VS_EXCEL = _SummaryPdfVsExcel(
    "Intersection Summary — TSMIS PDF vs Excel", "Intersection Summary",
    "TSMIS_PDF_vs_Excel_IntersectionSummary", IS_HEADER,
    lambda p: _rows_flat(p, _is.TSMIS_SHEET, IS_HEADER, "Intersection Summary"))
HIGHWAY_SUMMARY_PDF_VS_EXCEL = _SummaryPdfVsExcel(
    "Highway Summary — TSMIS PDF vs Excel", "Highway Summary",
    "TSMIS_PDF_vs_Excel_HighwaySummary", HS_HEADER,
    lambda p: _rows_flat(p, hsc.SHEET_NAME, HS_HEADER, "Highway Summary"))
