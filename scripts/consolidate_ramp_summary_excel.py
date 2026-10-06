"""Consolidate TSAR: Ramp Summary EXCEL exports into the same audited workbook
the PDF consolidator writes (v0.48.0).

The site's `rs_exportToExcel` writes one sheet ("Ramp Summary") that is the
print's two-column page laid out as one column of blocks: a title block
("TSAR - RAMPS SUMMARY" / "All Ramps on Route <r>" / "Reference Date: …" /
"Generated: …"), then the Highway Groups, On/Off Indicator, Population Groups
and Ramp Types blocks — each a section heading, a NUMBER | CODE header and
NUMBER | CODE rows — and the two totals ("Total Number of Ramps:" /
"Ramp Points w/out linework:") with the number beside the label.

`parse_xlsx` reads one file into EXACTLY the record `consolidate_ramp_summary
.parse_pdf` returns (the same schema regexes, the same CMP-AUD-019 matcher
diagnostics), so the shared pipeline writes the identical workbook and every
comparison — vs TSN, between environments, PDF vs Excel — reads either
edition. Censused on the 7/9, 7/23 and 10/2 statewide SSOR-prod pulls: all
126 routes of each parse field-for-field equal to the same day's print.

Console-free like the PDF consolidator; the pipeline lives there
(`consolidate_ramp_summary.consolidate_edition`).
"""
import re
from pathlib import Path

try:
    from openpyxl import load_workbook
    _DEPS_OK = True
except ImportError:
    _DEPS_OK = False

import consolidate_ramp_summary as rs
from paths import (OUTPUT_ROOT, output_day_dir, stamped_consolidated_filename)

SUBDIR = "ramp_summary_excel"
FILENAME = "tsar_ramp_summary_excel_consolidated.xlsx"
SHEET_NAME = "Ramp Summary"                  # the export's own sheet
REPORT_NAME = "Ramp Summary (Excel)"
INPUT_FMT = "Excel"

INPUT_DIR = OUTPUT_ROOT / SUBDIR
OUT_DIR = OUTPUT_ROOT / "consolidated"
OUT_PATH = OUT_DIR / FILENAME

_SECTIONS = {
    "Highway Groups": rs.HIGHWAY_GROUPS,
    "On/Off Indicator": rs.ONOFF,
    "Population Groups": rs.POP_GROUPS,
    "Ramp Types": rs.RAMP_TYPES,
}
_ROUTE_RE = re.compile(r"All Ramps on Route\s+(\d+\w*)$")
# The title-block lines the export writes above the first section.
_TITLE_LINES = (re.compile(r"TSAR\s*-\s*RAMPS SUMMARY$", re.I),
                re.compile(r"Reference Date:", re.I),
                re.compile(r"Generated:", re.I))
_TOTAL_LABEL = "Total Number of Ramps:"
_NOLW_LABEL = "Ramp Points w/out linework:"


def input_dir_for(day):
    """Per-route exports for `day` (a run-folder name); None = the legacy flat layout."""
    return (output_day_dir(day) / SUBDIR) if day else INPUT_DIR


def out_path_for(day):
    """The consolidated workbook for `day`; None = the legacy location."""
    if not day:
        return OUT_PATH
    return output_day_dir(day) / "consolidated" / stamped_consolidated_filename(FILENAME, day)


def _number(value, label, source):
    """A count cell as an int; anything else refuses the file (CMP-AUD-021 —
    a malformed count must never silently become a different number)."""
    if isinstance(value, bool):
        raise ValueError(f"{source}: the {label!r} count is a boolean")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value or "").strip().replace(",", "")
    if re.fullmatch(r"-?\d+", text):
        return int(text)
    raise ValueError(f"{source}: the {label!r} count {value!r} is not a whole number")


def parse_xlsx(path):
    """Return one flat dict for a Ramp Summary Excel export — the same keys,
    values and diagnostics `consolidate_ramp_summary.parse_pdf` produces for
    the print of the same route. A row that fits none of the censused shapes
    refuses the file (the export layout changed) instead of being dropped."""
    source = Path(path).name
    record = {"source_file": source, "route": None,
              "total_ramps": None, "ramp_points_no_linework": None}
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = [tuple(r[:2]) + (None,) * max(0, 2 - len(r))
                for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    blocks = {name: [] for name in _SECTIONS}
    section = None
    for row_no, (a, b) in enumerate(rows, start=1):
        a_text = "" if a is None else str(a).strip()
        b_text = "" if b is None else str(b).strip()
        if not a_text and not b_text:
            continue
        if isinstance(a, (int, float)) and not isinstance(a, bool):
            if section is None or not b_text:
                raise ValueError(f"{source}: row {row_no} carries a count outside any "
                                 "section — the export layout has changed")
            blocks[section].append((_number(a, b_text, source),
                                    re.sub(r"\s+", " ", b_text)))
            continue
        m = _ROUTE_RE.match(a_text)
        if m and not b_text:
            record["route"] = m.group(1)
            continue
        if a_text in _SECTIONS and not b_text:
            section = a_text
            continue
        if a_text == "NUMBER" and b_text == "CODE":
            continue
        if a_text == _TOTAL_LABEL:
            record["total_ramps"] = _number(b, a_text, source)
            section = None
            continue
        if a_text == _NOLW_LABEL:
            record["ramp_points_no_linework"] = _number(b, a_text, source)
            continue
        if section is None and not b_text and any(p.match(a_text) for p in _TITLE_LINES):
            continue
        raise ValueError(f"{source}: row {row_no} ({a_text!r}, {b_text!r}) is not "
                         "part of the censused Ramp Summary layout — the export "
                         "layout has changed; this app needs an update for it")
    unknown, duplicate = [], []
    for name, schema in _SECTIONS.items():
        used = set()
        record.update(rs.match_schema(blocks[name], schema, used))
        u, d = rs.schema_diagnostics(blocks[name], used, schema)
        unknown += u
        duplicate += d
    record["_unknown_rows"] = unknown
    record["_duplicate_rows"] = duplicate
    return record


EXCEL_EDITION = rs.Edition(
    glob="*.xlsx", parse=parse_xlsx, noun="file", report_name=REPORT_NAME,
    title="TSAR Ramp Summary (Excel) Consolidation", input_dir_for=input_dir_for,
    out_path_for=out_path_for, pdf_marker=False, missing_deps="openpyxl")


def consolidate(events=None, confirm_overwrite=None, day=None,
                input_dir=None, out_path=None, commit_guard=None):
    """Parse every per-route Ramp Summary Excel export into the audited
    workbook (the PDF consolidator's own layout). Console-free; returns a
    ConsolidateResult."""
    if not _DEPS_OK:
        from events import ConsolidateResult
        return ConsolidateResult(status="error",
                                 message="Required components are missing (openpyxl).")
    return rs.consolidate_edition(EXCEL_EDITION, events=events,
                                  confirm_overwrite=confirm_overwrite, day=day,
                                  input_dir=input_dir, out_path=out_path,
                                  commit_guard=commit_guard)


if __name__ == "__main__":
    from cli import run_consolidate_cli
    run_consolidate_cli(consolidate)
