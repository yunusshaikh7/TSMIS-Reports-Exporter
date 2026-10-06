"""Offline proof of the ArcGIS layer refresh (v0.46.0) — a stand-in ArcPy.

ArcPy only exists inside ArcGIS Pro, which neither the dev box nor CI has, so
the refresh is proven against a stand-in: `FAKE_ARCPY_SOURCE` is a tiny module
with exactly the ArcPy calls the worker makes (sign-in facts, MakeTableView,
GetCount, Table To Excel written with openpyxl, Delete, Describe). It runs two
ways from the one source:

  * in-process (`inprocess_launcher`) — the SHIPPED worker script is loaded
    from disk and run on a thread against the stand-in; the frozen exe's
    `--self-test` uses this, so a bundle missing the worker or breaking the
    refresh fails its own gate;
  * as a real `arcpy.py` beside a copy of the worker (build/check_arcgis_refresh
    drives the real subprocess launcher that way).

`run(tmp, emit)` is the scenario both callers share: a manual drop seeded in a
temp library, a refresh of three layers (two export, one is refused by the
"service"), every result checked — the library files, the backup, the rewritten
manifest read back with the builds' own reader, the drop identity — and then
the ArcGIS Pro check. Raises AssertionError on any failure.

A diagnostic driver: only self_test and the build checks import it.
"""
import json
import threading
import types
from datetime import datetime
from pathlib import Path

FAKE_ARCPY_SOURCE = r'''
"""A stand-in for ArcPy: just the calls the TSMIS layer export worker makes.
Configured by configure(spec) or by the JSON file named in TSMIS_FAKE_ARCPY:
{"version", "signed_in", "portal", "delay", "import_delay", "crash",
"layers": {id: {"name", "service_name", "header", "rows", "truncate",
"fail", "delay"}}}. `import_delay` stalls the import (a slow ArcGIS start),
`crash` kills the process at the first ArcPy call (a dying ArcGIS Python) and
a `delay` stalls Table To Excel (all layers, or one)."""
import json
import os
import time
from types import SimpleNamespace


class ExecuteError(Exception):
    pass


_spec = {}
_views = {}
_last = [""]
halt = [False]
env = SimpleNamespace(overwriteOutput=False)


def configure(spec):
    _spec.clear()
    _spec.update(spec)


if os.environ.get("TSMIS_FAKE_ARCPY"):
    with open(os.environ["TSMIS_FAKE_ARCPY"], encoding="utf-8") as _f:
        configure(json.load(_f))
    time.sleep(float(_spec.get("import_delay") or 0))


def _check_halt():
    if halt[0]:
        raise SystemExit("stopped")


def GetInstallInfo():
    if _spec.get("crash"):
        os._exit(9)
    return {"ProductName": "ArcGISPro", "Version": _spec.get("version", "3.3.5")}


def ProductInfo():
    return "ArcInfo"


def GetActivePortalURL():
    return _spec.get("portal", "https://rhapps-prod.dot.ca.gov/portal/")


def GetSigninToken():
    if _spec.get("signed_in"):
        return {"token": "fake-token", "referer": "fake-referer", "expires": 0}
    return None


def GetMessages(severity=0):
    return _last[0]


def _fail(message):
    _last[0] = message
    raise ExecuteError(message)


def _make_table_view(url, view):
    _check_halt()
    lid = str(url).rstrip("/").rsplit("/", 1)[-1]
    layer = (_spec.get("layers") or {}).get(lid)
    if layer is None:
        _fail("ERROR 000732: Input Rows: Dataset %s does not exist or is not supported" % url)
    if layer.get("fail"):
        _fail(layer["fail"])
    _views[view] = layer


def _get_count(view):
    _check_halt()
    return [str(len(_views[view]["rows"]))]


def _delete(view):
    _views.pop(view, None)


def _table_to_excel(view, out, alias="NAME", domains="CODE"):
    from openpyxl import Workbook
    layer = _views[view]
    end = time.time() + float(layer.get("delay") or _spec.get("delay") or 0)
    while time.time() < end:
        _check_halt()
        time.sleep(0.05)
    _check_halt()
    rows = layer["rows"]
    keep = len(rows) - int(layer.get("truncate") or 0)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(str(view)[:31])
    ws.append(list(layer["header"]))
    for row in rows[:keep]:
        ws.append(list(row))
    wb.save(out)


def Describe(view):
    layer = _views.get(view) or {}
    return SimpleNamespace(name=layer.get("service_name") or layer.get("name") or str(view))


management = SimpleNamespace(MakeTableView=_make_table_view, GetCount=_get_count,
                             Delete=_delete)
conversion = SimpleNamespace(TableToExcel=_table_to_excel)
'''

# The real SHS Tolls header as the in-app export writes it (the first work-PC
# refresh, 2026-10-05): 33 columns. The manual exports read the layer through a
# map and named the last column `Shape.STLength()`; read straight from the
# service it is `Shape__Length` (same values), so the first in-app refresh of a
# manual library reports exactly that one change on every line layer.
TOLLS_HEADER = [
    "OBJECTID", "District", "RouteNum", "RouteSuffix", "Alignment", "BeginCounty",
    "BeginPMPrefix", "BeginPMMeasure", "BeginPMSuffix", "EndCounty", "EndPMPrefix",
    "EndPMMeasure", "EndPMSuffix", "BeginODMeasure", "EndODMeasure", "Facility_Name",
    "InventoryItemStartDate", "InventoryItemEndDate", "RouteID", "FromARMeasure",
    "ToARMeasure", "LRSFromDate", "LRSToDate", "EventID", "CreatedUser",
    "LastEditedUser", "CreatedDate", "LastEditedDate", "ARSStatus", "LocError",
    "GlobalID", "Toll_Type", "Shape__Length"]
MANUAL_TOLLS_HEADER = TOLLS_HEADER[:-1] + ["Shape.STLength()"]   # the 2026-08-19 drop
TRANSITION_NOTE = "columns changed (+Shape__Length) (−Shape.STLength())"
CITY_HEADER = ["OBJECTID", "District", "RouteNum", "BeginCounty", "City_Code",
               "RouteID", "LRSFromDate", "LRSToDate"]


def _tolls_rows(n):
    return [[i, "District 12", "073", ".", "Right", "Orange", ".", round(10 + i / 10, 3),
             ".", "Orange", ".", 23.75, ".", 0.1, 13.75, "Route 73", None, None,
             "SHS_073._P", 0.09, 13.78, "2019-10-14", None, f"EV{i:04d}", "RH", "SDE",
             "2025-09-23", "2025-09-23", None, "NO ERROR", f"GID{i:04d}", "Toll Roads",
             22000.61] for i in range(1, n + 1)]


def _city_rows(n):
    return [[i, "District 7", "001", "Los Angeles", "LB", "SHS_001_P", "2020-01-01", None]
            for i in range(1, n + 1)]


def sample_spec(**over):
    """The stand-in service: SHS Tolls (27 rows) and City (40) export; Route
    Direction is refused the way an unreachable layer is."""
    spec = {"version": "3.3.5", "signed_in": False, "delay": 0,
            "layers": {"138": {"name": "SHS Tolls", "header": TOLLS_HEADER,
                               "rows": _tolls_rows(27)},
                       "74": {"name": "City", "header": CITY_HEADER,
                              "rows": _city_rows(40)}}}
    spec.update(over)
    return spec


def make_fake_arcpy(spec):
    """The stand-in as a live module object, configured with `spec`."""
    mod = types.ModuleType("arcpy")
    exec(compile(FAKE_ARCPY_SOURCE, "<fake arcpy>", "exec"), mod.__dict__)
    mod.configure(spec)
    return mod


class _ThreadHandle:
    """A launcher handle for the in-process worker: poll() gives the exit code
    once the thread ends; stop() asks the stand-in to halt."""

    def __init__(self, target, fake):
        self._fake = fake
        self._code = None
        self._thread = threading.Thread(target=self._run, args=(target,),
                                        name="arcgis-selftest-worker", daemon=True)
        self._thread.start()

    def _run(self, target):
        try:
            self._code = int(target() or 0)
        except BaseException:      # noqa: BLE001  # silent-ok: a halted or crashed stand-in worker is exit code 1, as a killed process would be
            self._code = 1

    def poll(self):
        return None if self._thread.is_alive() else self._code

    def stop(self):
        if self._thread.is_alive():
            self._fake.halt[0] = True
        self._thread.join(timeout=10)

    def close(self):
        self.stop()


def inprocess_launcher(fake):
    """A launcher that runs the SHIPPED worker script on a thread against
    `fake`, writing the same events file the real process would."""
    import runpy

    def launch(python, script, request_path, events_path, log_path):
        ns = runpy.run_path(str(script))
        request = json.loads(Path(request_path).read_text(encoding="utf-8"))
        emit = ns["Emitter"](str(events_path))
        Path(log_path).write_text("in-process worker\n", encoding="utf-8")

        def target():
            try:
                return ns["run"](request, emit, arcpy=fake)
            finally:
                emit.close()

        return _ThreadHandle(target, fake)

    return launch


def seed_manual_library(lib):
    """A manual drop in `lib`: a six-column manifest (stamped 2026-08-19) and an
    older three-row SHS Tolls file with the manual export's column names."""
    from openpyxl import Workbook

    import clean_road_layers as crl

    lib = Path(lib)
    lib.mkdir(parents=True, exist_ok=True)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("SHS Tolls")
    ws.append(MANUAL_TOLLS_HEADER)
    for row in _tolls_rows(3):
        ws.append(row)
    wb.save(lib / "11_SHS Tolls.xlsx")
    idx = Workbook(write_only=True)
    sheet = idx.create_sheet("INDEX")
    sheet.append(crl.INDEX_HEADER)
    sheet.append(["11_SHS Tolls.xlsx", "SHS Tolls", 3, len(MANUAL_TOLLS_HEADER), "SHS Tolls",
                  "https://example.invalid/FeatureServer;VERSION=sde.DEFAULT/138"])
    idx.properties.created = datetime(2026, 8, 19, 19, 14, 14)
    idx.save(lib / crl.INDEX_NAME)


def _check(cond, what):
    if not cond:
        raise AssertionError(f"arcgis layer refresh: {what}")


def run(tmp, emit=print):
    """The shared scenario (see the module docstring). Returns a one-line
    summary; raises AssertionError on any failure."""
    import arcgis_layers
    import arcgis_refresh as ar
    import clean_road_layers as crl
    import paths

    lib = Path(tmp) / "arcgis_layers"
    seed_manual_library(lib)
    saved_root = paths.ARCGIS_LAYERS_ROOT
    paths.ARCGIS_LAYERS_ROOT = lib
    try:
        fake = make_fake_arcpy(sample_spec())
        logs = []
        events = _events(logs)
        res = ar.refresh(["SHS Tolls", "City", "Route Direction"], events=events,
                         python="fake-python.exe", launcher=inprocess_launcher(fake))
        by = {r["name"]: r for r in res.layers}
        _check(res.status == "partial", f"status {res.status!r}: {res.message}")
        _check(by["SHS Tolls"]["status"] == "exported" and by["SHS Tolls"]["rows"] == 27,
               f"SHS Tolls {by['SHS Tolls']}")
        _check(by["SHS Tolls"].get("message") == TRANSITION_NOTE,
               f"the manual-to-in-app column note {by['SHS Tolls'].get('message')!r}")
        _check(by["City"]["status"] == "exported" and by["City"]["rows"] == 40,
               f"City {by['City']}")
        _check(by["Route Direction"]["status"] == "failed"
               and "000732" in (by["Route Direction"].get("message") or ""),
               f"Route Direction {by['Route Direction']}")
        _check((lib / "11_SHS Tolls.xlsx").is_file(), "SHS Tolls kept its library name")
        _check((lib / "33_City.xlsx").is_file(), "City took its library number")
        _check((lib / "_previous" / "11_SHS Tolls.xlsx").is_file()
               and (lib / "_previous" / crl.INDEX_NAME).is_file(),
               "the replaced file and manifest are in _previous")
        _check(not any(ar.staging_dir().glob("*")), "staging is empty")
        index = crl.read_index(lib)                    # the builds' own reader
        _check(index["SHS Tolls"]["rows"] == 27 and index["City"]["rows"] == 40,
               f"manifest rows {index}")
        n = sum(1 for _r in crl.stream_layer(lib / "11_SHS Tolls.xlsx",
                                             ["OBJECTID", "District", "LRSFromDate"],
                                             layer_name="SHS Tolls",
                                             expected_rows=index["SHS Tolls"]["rows"]))
        _check(n == 27, f"the builds' reader saw {n} SHS Tolls rows")
        full, _created = arcgis_layers.read_index_full(lib)
        _check(all(full[k]["exported_at"] for k in ("SHS Tolls", "City")),
               "every refreshed layer records its export time")
        drop = arcgis_layers.drop_info(lib)
        _check(drop["fingerprint"] and drop["exported"] == datetime.now().date().isoformat()
               and set(drop["layers"]) == {"SHS Tolls", "City"}, f"drop {drop}")
        run_rec = ar.last_run() or {}
        _check(run_rec.get("status") == "partial", f"last_run {run_rec.get('status')!r}")

        probe = ar.probe(events=events, python="fake-python.exe",
                         launcher=inprocess_launcher(make_fake_arcpy(sample_spec())))
        _check(probe.get("ok") and probe.get("rows") == 27
               and probe.get("library_match") is True and probe.get("dialect") == "labels",
               f"probe {probe}")
    finally:
        paths.ARCGIS_LAYERS_ROOT = saved_root
    summary = (f"layer refresh: 2 exported + 1 refused, swapped/backed up/indexed; "
               f"builds' reader ok; ArcGIS Pro check ok ({len(logs)} log lines)")
    emit(summary)
    return summary


def _events(sink):
    from events import Events
    return Events(on_log=sink.append)
